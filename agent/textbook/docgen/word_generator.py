import os
import re
from typing import List, Optional
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from .style_template import StyleTemplate, get_template

class MarkdownToWordConverter:
    BODY_EAST_ASIA_FONT = "仿宋"
    BODY_SIZE = 12
    HEADING_EAST_ASIA_FONT = "黑体"
    HEADING_SIZE = 16
    CAPTION_EAST_ASIA_FONT = "仿宋"
    NOTE_EAST_ASIA_FONT = "宋体"
    SMALL_SIZE = 10.5

    def __init__(self, template: StyleTemplate = None):
        self.template = template or get_template("academic")
        self.doc = Document()
        self._setup_page()
        self._setup_styles()

    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(self.template.margin_top)
        section.bottom_margin = Cm(self.template.margin_bottom)
        section.left_margin = Cm(self.template.margin_left)
        section.right_margin = Cm(self.template.margin_right)

    def _setup_styles(self):
        for level, (font, size, bold, align) in enumerate([
            (self.HEADING_EAST_ASIA_FONT, self.HEADING_SIZE, True, self.template.heading1_alignment),
            (self.HEADING_EAST_ASIA_FONT, self.HEADING_SIZE, True, "left"),
            (self.HEADING_EAST_ASIA_FONT, self.HEADING_SIZE, False, "left"),
        ], 1):
            style_name = f'Heading {level}'
            try:
                style = self.doc.styles[style_name]
            except KeyError:
                style = self.doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = font
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_before = Pt(size * 0.5)
            style.paragraph_format.space_after = Pt(size * 0.5)
            if align == "center":
                style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rpr = style.element.get_or_add_rPr()
            rFonts = rpr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = rpr.makeelement(qn('w:rFonts'), {})
                rpr.insert(0, rFonts)
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            rFonts.set(qn('w:eastAsia'), font)

        normal = self.doc.styles['Normal']
        normal.font.name = self.BODY_EAST_ASIA_FONT
        normal.font.size = Pt(self.BODY_SIZE)
        normal.paragraph_format.line_spacing = 1.5
        rpr = normal.element.get_or_add_rPr()
        rFonts = rpr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = rpr.makeelement(qn('w:rFonts'), {})
            rpr.insert(0, rFonts)
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), self.BODY_EAST_ASIA_FONT)

    @staticmethod
    def _set_run_font(run, east_asia: str, size: float, *, ascii_font: str = "Times New Roman", bold=None, italic=None):
        run.font.name = ascii_font
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        rpr = run._element.get_or_add_rPr()
        rFonts = rpr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = rpr.makeelement(qn('w:rFonts'), {})
            rpr.insert(0, rFonts)
        rFonts.set(qn('w:ascii'), ascii_font)
        rFonts.set(qn('w:hAnsi'), ascii_font)
        rFonts.set(qn('w:eastAsia'), east_asia)

    @staticmethod
    def _set_paragraph_spacing(paragraph, *, line_spacing=1.5, before=0, after=0, first_line_indent=None, alignment=None):
        paragraph.paragraph_format.line_spacing = line_spacing
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        if first_line_indent is not None:
            paragraph.paragraph_format.first_line_indent = first_line_indent
        if alignment is not None:
            paragraph.paragraph_format.alignment = alignment

    def add_title(self, title: str):
        p = self.doc.add_heading(title, level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            self._set_run_font(run, self.HEADING_EAST_ASIA_FONT, self.HEADING_SIZE, bold=True)
        self._set_paragraph_spacing(p, line_spacing=1.5, before=8, after=8)

    def add_heading(self, text: str, level: int):
        if level < 1 or level > 3:
            level = 3
        heading = self.doc.add_heading('', level=level)
        self._add_formatted_runs(heading, text)
        east_asia = self.HEADING_EAST_ASIA_FONT
        bold = level in (1, 2)
        for run in heading.runs:
            self._set_run_font(run, east_asia, self.HEADING_SIZE, bold=bold)
        self._set_paragraph_spacing(heading, line_spacing=1.5, before=8, after=8)

    def add_paragraph(self, text: str):
        p = self.doc.add_paragraph()
        self._add_formatted_runs(p, text)
        self._style_body_paragraph(p)
        return p

    def _style_body_paragraph(self, paragraph):
        self._set_paragraph_spacing(
            paragraph,
            line_spacing=1.5,
            first_line_indent=Pt(self.BODY_SIZE * self.template.body_first_line_indent),
        )
        for run in paragraph.runs:
            self._set_run_font(run, self.BODY_EAST_ASIA_FONT, self.BODY_SIZE)

    def _add_formatted_runs(self, paragraph, text: str):
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`[^`]+`)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                self._set_run_font(run, self.BODY_EAST_ASIA_FONT, self.BODY_SIZE, bold=True)
            elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
                run = paragraph.add_run(part[1:-1])
                self._set_run_font(run, self.BODY_EAST_ASIA_FONT, self.BODY_SIZE, italic=True)
            elif part.startswith('`') and part.endswith('`'):
                run = paragraph.add_run(part[1:-1])
                self._set_run_font(run, self.BODY_EAST_ASIA_FONT, self.BODY_SIZE)
            else:
                run = paragraph.add_run(part)
                self._set_run_font(run, self.BODY_EAST_ASIA_FONT, self.BODY_SIZE)

    def add_code_block(self, code: str, language: str = ""):
        p = self.doc.add_paragraph()
        run = p.add_run(code)
        self._set_run_font(run, self.BODY_EAST_ASIA_FONT, self.BODY_SIZE)
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5

    def add_image(self, image_path: str, caption: str = "", width: float = None):
        if os.path.exists(image_path):
            if width:
                self.doc.add_picture(image_path, width=Inches(width))
            else:
                self.doc.add_picture(image_path, width=Inches(5.0))
            if caption:
                p = self.doc.add_paragraph(caption)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._set_paragraph_spacing(p, line_spacing=1.5, after=5.25, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                for run in p.runs:
                    self._set_run_font(run, self.CAPTION_EAST_ASIA_FONT, self.SMALL_SIZE)

    def add_table(self, headers: List[str], rows: List[List[str]]):
        table = self.doc.add_table(rows=len(rows)+1, cols=len(headers))
        table.style = 'Table Grid'
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                table.rows[r+1].cells[c].text = val
        self._style_table(table)

    def _style_table(self, table):
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self._set_paragraph_spacing(
                        paragraph,
                        line_spacing=1.0,
                        first_line_indent=Pt(0),
                        alignment=WD_ALIGN_PARAGRAPH.CENTER,
                    )
                    for run in paragraph.runs:
                        self._set_run_font(run, self.CAPTION_EAST_ASIA_FONT, self.SMALL_SIZE)
                        if row_idx == 0:
                            run.bold = True

    def add_bullet_list(self, items: List[str]):
        for item in items:
            p = self.doc.add_paragraph(style='List Bullet')
            self._add_formatted_runs(p, item)
            self._style_body_paragraph(p)

    def add_numbered_list(self, items: List[str]):
        for item in items:
            p = self.doc.add_paragraph(style='List Number')
            self._add_formatted_runs(p, item)
            self._style_body_paragraph(p)

    def add_quote(self, text: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        self._add_formatted_runs(p, text)
        self._style_body_paragraph(p)
        for run in p.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(100, 100, 100)

    def add_caption(self, text: str, kind: str = "figure"):
        p = self.doc.add_paragraph()
        p.add_run(text)
        before = 5.25 if kind == "table" else 0
        after = 5.25 if kind == "figure" else 0
        self._set_paragraph_spacing(p, line_spacing=1.5, before=before, after=after, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        for run in p.runs:
            self._set_run_font(run, self.CAPTION_EAST_ASIA_FONT, self.SMALL_SIZE)
        return p

    def add_note(self, text: str):
        p = self.doc.add_paragraph()
        p.add_run(text)
        self._set_paragraph_spacing(p, line_spacing=1.5)
        for run in p.runs:
            self._set_run_font(run, self.NOTE_EAST_ASIA_FONT, self.SMALL_SIZE)
        return p

    def convert_markdown(self, markdown_text: str):
        lines = markdown_text.split('\n')
        i = 0
        in_code_block = False
        code_buffer = []
        code_lang = ""

        while i < len(lines):
            line = lines[i]

            if line.strip().startswith('```'):
                if in_code_block:
                    self.add_code_block('\n'.join(code_buffer), code_lang)
                    code_buffer = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lang = line.strip()[3:].strip()
                i += 1
                continue

            if in_code_block:
                code_buffer.append(line)
                i += 1
                continue

            stripped = line.strip()

            if not stripped:
                i += 1
                continue

            if stripped.startswith('#'):
                level = len(stripped) - len(stripped.lstrip('#'))
                text = stripped.lstrip('#').strip()
                if level <= 3:
                    self.add_heading(text, level)
                else:
                    self.add_paragraph(text)
            elif stripped.startswith('> '):
                self.add_quote(stripped[2:])
            elif stripped.startswith('- ') or stripped.startswith('* '):
                items = []
                while i < len(lines) and (lines[i].strip().startswith('- ') or lines[i].strip().startswith('* ')):
                    items.append(lines[i].strip()[2:])
                    i += 1
                self.add_bullet_list(items)
                continue
            elif re.match(r'^\d+\.\s', stripped):
                items = []
                while i < len(lines) and re.match(r'^\d+\.\s', lines[i].strip()):
                    items.append(re.sub(r'^\d+\.\s', '', lines[i].strip()))
                    i += 1
                self.add_numbered_list(items)
                continue
            elif stripped.startswith('|') and '|' in stripped[1:]:
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    table_lines.append(lines[i].strip())
                    i += 1
                if len(table_lines) >= 2:
                    headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                    rows = []
                    for tl in table_lines[2:]:
                        row = [c.strip() for c in tl.split('|')[1:-1]]
                        rows.append(row)
                    self.add_table(headers, rows)
                continue
            elif stripped.startswith('!['):
                match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
                if match:
                    caption = match.group(1)
                    path = match.group(2)
                    self.add_image(path, caption)
            elif re.match(r'^(表|Table)\s*[\d一二三四五六七八九十]+', stripped, re.IGNORECASE):
                self.add_caption(stripped, kind="table")
            elif re.match(r'^(图|Figure)\s*[\d一二三四五六七八九十]+', stripped, re.IGNORECASE):
                self.add_caption(stripped, kind="figure")
            elif re.match(r'^注[:：]', stripped):
                self.add_note(stripped)
            else:
                self.add_paragraph(stripped)

            i += 1

    def save(self, filepath: str):
        self.doc.save(filepath)

    @classmethod
    def from_markdown(cls, markdown_text: str, template_name: str = "academic", output_path: str = "") -> str:
        converter = cls(get_template(template_name))
        converter.convert_markdown(markdown_text)
        if output_path:
            converter.save(output_path)
        return output_path
