import os
import re
import json
import shutil
import subprocess
import sys
import time
import hashlib
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Optional
from urllib.parse import urljoin

from common.log import logger
from config import conf


class KnowledgeService:

    def __init__(self, workspace_root: str, on_progress=None):
        self.workspace_root = workspace_root
        self.knowledge_dir = os.path.join(workspace_root, "knowledge")
        self.on_progress = on_progress

    def _emit_progress(self, stage: str, message: str = "", **extra):
        if not self.on_progress:
            return
        try:
            payload = {"stage": stage, "message": message, **extra}
            self.on_progress(payload)
        except Exception:
            pass

    def _resolve_book_dir(self, book_id: str = "") -> str:
        if book_id and (".." in book_id or "/" in book_id or "\\" in book_id):
            raise ValueError("invalid book_id")
        if book_id:
            base = os.path.join(self.knowledge_dir, book_id)
        else:
            base = self.knowledge_dir
        base = os.path.normpath(base)
        knowledge_dir_norm = os.path.normpath(self.knowledge_dir)
        if not base.startswith(knowledge_dir_norm + os.sep) and base != knowledge_dir_norm:
            raise ValueError("path outside knowledge dir")
        return base

    def list_tree(self, book_id: str = "") -> dict:
        base = self._resolve_book_dir(book_id)
        if not os.path.isdir(base):
            return {"tree": [], "stats": {"pages": 0, "size": 0}, "enabled": conf().get("knowledge", True)}
        stats = {"pages": 0, "size": 0}
        root_files, tree = self._scan_dir(base, stats, base_dir=base, is_root=True)
        return {
            "root_files": root_files,
            "tree": tree,
            "stats": stats,
            "enabled": conf().get("knowledge", True),
        }

    def list_files_page(
        self,
        book_id: str = "",
        offset: int = 0,
        limit: int = 80,
        path_prefix: str = "",
        query: str = "",
    ) -> dict:
        base = self._resolve_book_dir(book_id)
        if not os.path.isdir(base):
            return {"files": [], "offset": 0, "limit": limit, "total": 0, "has_more": False}
        offset = max(0, int(offset or 0))
        limit = min(300, max(1, int(limit or 80)))
        query = (query or "").strip().lower()
        path_prefix = (path_prefix or "").strip().strip("/\\")
        files = []
        total_size = 0
        allowed_exts = {".md", ".txt", ".json", ".csv"}
        for root, dirs, names in os.walk(base):
            dirs[:] = sorted(d for d in dirs if not d.startswith("."))
            rel_dir = os.path.relpath(root, base).replace("\\", "/")
            if rel_dir == ".":
                rel_dir = ""
            if path_prefix and rel_dir and not rel_dir.startswith(path_prefix):
                continue
            for name in sorted(names):
                if name.startswith(".") or os.path.splitext(name)[1].lower() not in allowed_exts:
                    continue
                full = os.path.join(root, name)
                rel = os.path.relpath(full, base).replace("\\", "/")
                if path_prefix and not rel.startswith(path_prefix):
                    continue
                if query and query not in rel.lower() and query not in name.lower():
                    continue
                size = os.path.getsize(full)
                total_size += size
                files.append({
                    "name": name,
                    "title": name.replace(".md", ""),
                    "path": rel,
                    "dir": rel.rsplit("/", 1)[0] if "/" in rel else "root",
                    "size": size,
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S", time.localtime(os.path.getmtime(full))),
                })
        total = len(files)
        page = files[offset:offset + limit]
        return {
            "files": page,
            "offset": offset,
            "limit": limit,
            "total": total,
            "total_size": total_size,
            "has_more": offset + limit < total,
            "next_offset": offset + len(page),
        }

    def _scan_dir(self, dir_path: str, stats: dict, base_dir: str, is_root: bool = False) -> tuple:
        files = []
        children = []
        for name in sorted(os.listdir(dir_path)):
            if name.startswith("."):
                continue
            full = os.path.join(dir_path, name)
            if os.path.isdir(full):
                sub_files, sub_children = self._scan_dir(full, stats, base_dir)
                children.append({
                    "dir": name,
                    "path": os.path.relpath(full, base_dir).replace("\\", "/"),
                    "files": sub_files,
                    "children": sub_children,
                })
            elif os.path.splitext(name)[1].lower() in {".md", ".txt", ".json", ".csv"}:
                size = os.path.getsize(full)
                if not is_root:
                    stats["pages"] += 1
                    stats["size"] += size
                title = name.replace(".md", "")
                try:
                    with open(full, "r", encoding="utf-8") as f:
                        first_line = f.readline().strip()
                    if first_line.startswith("# "):
                        title = first_line[2:].strip()
                except Exception:
                    pass
                files.append({
                    "name": name,
                    "title": title,
                    "path": os.path.relpath(full, base_dir).replace("\\", "/"),
                    "size": size,
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S", time.localtime(os.path.getmtime(full))),
                })
        return files, children

    def read_file(self, rel_path: str, book_id: str = "") -> dict:
        if not rel_path or ".." in rel_path:
            raise ValueError("invalid path")
        base = self._resolve_book_dir(book_id)
        full_path = os.path.normpath(os.path.join(base, rel_path))
        allowed = os.path.normpath(base)
        if not full_path.startswith(allowed + os.sep) and full_path != allowed:
            raise ValueError("path outside knowledge dir")
        if not os.path.isfile(full_path):
            raise FileNotFoundError(f"file not found: {rel_path}")
        with open(full_path, "r", encoding="utf-8") as f:
            content = f.read(300000)
            truncated = bool(f.read(1))
        return {"content": content, "path": rel_path, "truncated": truncated, "max_chars": 300000}

    def build_graph(
        self,
        book_id: str = "",
        limit: int = 140,
        focus_id: str = "",
        depth: int = 1,
        query: str = "",
        min_confidence: float = 0.0,
    ) -> dict:
        base = self._resolve_book_dir(book_id)
        knowledge_path = Path(base)
        if not knowledge_path.is_dir():
            return {"nodes": [], "edges": [], "links": []}
        wiki_graph = os.path.join(self._wiki_base_dir(book_id), "graph.json")
        if os.path.isfile(wiki_graph):
            try:
                with open(wiki_graph, "r", encoding="utf-8") as f:
                    graph = json.load(f)
                edges = graph.get("edges") or graph.get("links") or []
                return self._project_graph(
                    graph.get("nodes", []),
                    edges,
                    limit=limit,
                    focus_id=focus_id,
                    depth=depth,
                    query=query,
                    min_confidence=min_confidence,
                    include_links=True,
                )
            except Exception:
                pass
        nodes = {}
        links = []
        link_re = re.compile(r'\[([^\]]*)\]\(([^)]+\.md)\)')
        for md_file in knowledge_path.rglob("*.md"):
            rel = str(md_file.relative_to(knowledge_path))
            if rel in ("index.md", "log.md"):
                continue
            parts = rel.replace("\\", "/").split("/")
            category = parts[0] if len(parts) > 1 else "root"
            title = md_file.stem.replace("-", " ").title()
            try:
                content = md_file.read_text(encoding="utf-8")
                first_line = content.strip().split("\n")[0]
                if first_line.startswith("# "):
                    title = first_line[2:].strip()
                for _, link_target in link_re.findall(content):
                    resolved = (md_file.parent / link_target).resolve()
                    try:
                        target_rel = str(resolved.relative_to(knowledge_path))
                    except ValueError:
                        continue
                    if target_rel != rel:
                        links.append({"source": rel, "target": target_rel})
            except Exception:
                pass
            nodes[rel] = {"id": rel, "label": title, "category": category}
        valid_ids = set(nodes.keys())
        links = [l for l in links if l["source"] in valid_ids and l["target"] in valid_ids]
        seen = set()
        deduped = []
        for l in links:
            key = tuple(sorted([l["source"], l["target"]]))
            if key not in seen:
                seen.add(key)
                deduped.append(l)
        return self._project_graph(
            list(nodes.values()),
            deduped,
            limit=limit,
            focus_id=focus_id,
            depth=depth,
            query=query,
            min_confidence=min_confidence,
            include_links=True,
        )

    def upload_document(self, source_path: str, category: str = "sources", book_id: str = "") -> dict:
        if not source_path or ".." in source_path:
            raise ValueError("invalid source path")
        if not os.path.isfile(source_path):
            raise FileNotFoundError(f"source file not found: {source_path}")
        if category and (".." in category or "/" in category or "\\" in category):
            raise ValueError("invalid category")
        if book_id and (".." in book_id or "/" in book_id or "\\" in book_id):
            raise ValueError("invalid book_id")
        if book_id:
            target_dir = os.path.join(self.knowledge_dir, book_id, category)
        else:
            target_dir = os.path.join(self.knowledge_dir, category)
        target_dir = os.path.normpath(target_dir)
        knowledge_dir_norm = os.path.normpath(self.knowledge_dir)
        if not target_dir.startswith(knowledge_dir_norm + os.sep) and target_dir != knowledge_dir_norm:
            raise ValueError("target path outside knowledge dir")
        os.makedirs(target_dir, exist_ok=True)
        file_name = os.path.basename(source_path)
        dest_path = os.path.join(target_dir, file_name)
        shutil.copy2(source_path, dest_path)
        file_size = os.path.getsize(dest_path)
        rel_path = os.path.relpath(dest_path, self.knowledge_dir).replace("\\", "/")
        return {"name": file_name, "category": category, "path": rel_path, "size": file_size, "status": "ready"}

    def save_web_source(
        self,
        url: str,
        title: str,
        content: str,
        reason: str = "",
        book_id: str = "",
        tags: Optional[list] = None,
        force: bool = False,
    ) -> dict:
        """Save a useful web source as Markdown in the knowledge sources directory."""
        url = (url or "").strip()
        title = (title or "").strip() or "Web source"
        content = self._normalize_web_source_content(content)
        reason = (reason or "").strip()
        tags = [str(tag).strip() for tag in (tags or []) if str(tag).strip()]

        useful, rejection_reason = self._is_useful_web_source(url, content, reason, force=force)
        if not useful:
            return {"useful": False, "reason": rejection_reason, "status": "skipped"}

        target_dir = os.path.join(self._resolve_book_dir(book_id), "sources")
        target_dir = os.path.normpath(target_dir)
        knowledge_dir_norm = os.path.normpath(self.knowledge_dir)
        if not target_dir.startswith(knowledge_dir_norm + os.sep) and target_dir != knowledge_dir_norm:
            raise ValueError("target path outside knowledge dir")
        os.makedirs(target_dir, exist_ok=True)

        content_hash = hashlib.sha256(f"{url}\n{content}".encode("utf-8", errors="ignore")).hexdigest()[:12]
        slug = self._safe_slug(title, "web")
        file_name = f"web_{slug}_{content_hash}.md"
        dest_path = os.path.join(target_dir, file_name)
        if os.path.exists(dest_path):
            rel_path = os.path.relpath(dest_path, self.knowledge_dir).replace("\\", "/")
            return {
                "useful": True,
                "status": "exists",
                "name": file_name,
                "path": rel_path,
                "size": os.path.getsize(dest_path),
            }

        captured_at = time.strftime("%Y-%m-%dT%H:%M:%S")
        markdown = self._format_web_source_markdown(
            title=title,
            url=url,
            content=content,
            reason=reason,
            tags=tags,
            captured_at=captured_at,
            content_hash=content_hash,
        )
        with open(dest_path, "w", encoding="utf-8") as f:
            f.write(markdown)

        self.sync_compat_index(book_id)
        rel_path = os.path.relpath(dest_path, self.knowledge_dir).replace("\\", "/")
        return {
            "useful": True,
            "status": "saved",
            "name": file_name,
            "path": rel_path,
            "size": os.path.getsize(dest_path),
            "content_hash": content_hash,
        }

    def _normalize_web_source_content(self, content: str) -> str:
        content = (content or "").replace("\r\n", "\n").replace("\r", "\n")
        content = re.sub(r"\n{3,}", "\n\n", content)
        return "\n".join(line.rstrip() for line in content.split("\n")).strip()

    def _is_useful_web_source(self, url: str, content: str, reason: str, force: bool = False) -> tuple:
        if force:
            return True, ""
        if not url.startswith(("http://", "https://")):
            return False, "URL must be http or https."
        if len(content) < 180:
            return False, "Content is too short to be a useful knowledge source."
        try:
            from urllib.parse import urlparse
            host = urlparse(url).netloc.lower()
        except Exception:
            host = ""
        if any(search_host in host for search_host in ("google.", "baidu.", "bing.", "duckduckgo.", "search.brave.")):
            return False, "Search result pages should not be stored as knowledge sources; fetch the original result URL."
        if len(reason) < 12:
            return False, "A usefulness reason is required so future agents know when to use this source."
        return True, ""

    def _format_web_source_markdown(
        self,
        title: str,
        url: str,
        content: str,
        reason: str,
        tags: list,
        captured_at: str,
        content_hash: str,
    ) -> str:
        safe_title = title.replace("\n", " ").strip()
        tag_lines = "\n".join(f"  - {tag}" for tag in tags) if tags else "  - web"
        return (
            "---\n"
            "source_type: web\n"
            f"title: {json.dumps(safe_title, ensure_ascii=False)}\n"
            f"url: {json.dumps(url, ensure_ascii=False)}\n"
            f"captured_at: {captured_at}\n"
            f"content_hash: {content_hash}\n"
            "tags:\n"
            f"{tag_lines}\n"
            "---\n\n"
            f"# {safe_title}\n\n"
            "## Source\n\n"
            f"- URL: {url}\n"
            f"- Captured at: {captured_at}\n\n"
            "## Use When\n\n"
            f"{reason or 'Use this source when its topic matches the current textbook section.'}\n\n"
            "## Extracted Content\n\n"
            f"{content}\n"
        )

    def list_sources(self, book_id: str = "") -> dict:
        if book_id and (".." in book_id or "/" in book_id or "\\" in book_id):
            raise ValueError("invalid book_id")
        if book_id:
            scan_dir = os.path.join(self.knowledge_dir, book_id)
        else:
            scan_dir = self.knowledge_dir
        scan_dir = os.path.normpath(scan_dir)
        knowledge_dir_norm = os.path.normpath(self.knowledge_dir)
        if not scan_dir.startswith(knowledge_dir_norm + os.sep) and scan_dir != knowledge_dir_norm:
            raise ValueError("scan path outside knowledge dir")
        sources = []
        if not os.path.isdir(scan_dir):
            return {"sources": sources, "total": 0}
        for root, dirs, files in os.walk(scan_dir):
            dirs[:] = [d for d in dirs if not d.startswith(".") and d not in {"_llm_wiki", "_parsed", "_raw"}]
            for fname in sorted(files):
                if fname.startswith("."):
                    continue
                full_path = os.path.join(root, fname)
                try:
                    stat = os.stat(full_path)
                except OSError:
                    continue
                rel = os.path.relpath(full_path, self.knowledge_dir).replace("\\", "/")
                parts = rel.split("/")
                if book_id and len(parts) > 2:
                    cat = parts[1]
                elif len(parts) > 1:
                    cat = parts[0]
                else:
                    cat = "root"
                sources.append({
                    "name": fname,
                    "category": cat,
                    "size": stat.st_size,
                    "status": "ready",
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S", time.localtime(stat.st_mtime)),
                })
        return {"sources": sources, "total": len(sources)}

    def get_status(self, book_id: str = "") -> dict:
        base = self._resolve_book_dir(book_id)
        if not os.path.isdir(base):
            return {"total_documents": 0, "total_size": 0, "categories": [], "enabled": conf().get("knowledge", True)}
        self.sync_compat_index(book_id)
        wiki_counts = self._aggregate_wiki_counts(book_id)
        total_docs = 0
        total_size = 0
        categories = set()
        for root, dirs, files in os.walk(base):
            dirs[:] = [d for d in dirs if not d.startswith(".") and d not in {"_llm_wiki", "_parsed", "_raw"}]
            for fname in files:
                if fname.startswith("."):
                    continue
                full_path = os.path.join(root, fname)
                try:
                    stat = os.stat(full_path)
                except OSError:
                    continue
                total_docs += 1
                total_size += stat.st_size
                rel = os.path.relpath(full_path, base).replace("\\", "/")
                parts = rel.split("/")
                if len(parts) > 1:
                    categories.add(parts[0])
        return {
            "total_documents": total_docs,
            "total_size": total_size,
            "categories": sorted(categories),
            "enabled": conf().get("knowledge", True),
            "wiki": wiki_counts,
        }

    def _aggregate_wiki_counts(self, book_id: str = "") -> dict:
        if book_id:
            return self._index_counts(self._load_wiki_index(book_id))

        totals = {"sources": 0, "chunks": 0, "pages": 0, "entities": 0, "relations": 0}
        seen_indexes = set()
        root_index = os.path.join(self._wiki_base_dir(""), "index.json")
        candidate_indexes = []
        if os.path.isfile(root_index):
            candidate_indexes.append(root_index)
        if os.path.isdir(self.knowledge_dir):
            for child in sorted(os.listdir(self.knowledge_dir)):
                child_index = os.path.join(self.knowledge_dir, child, "_llm_wiki", "index.json")
                if os.path.isfile(child_index):
                    candidate_indexes.append(child_index)
        for index_path in candidate_indexes:
            norm = os.path.normcase(os.path.abspath(index_path))
            if norm in seen_indexes:
                continue
            seen_indexes.add(norm)
            try:
                with open(index_path, "r", encoding="utf-8") as f:
                    counts = self._index_counts(json.load(f))
                for key in totals:
                    totals[key] += int(counts.get(key, 0) or 0)
            except Exception:
                continue
        return totals

    def _extract_text(self, file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return self._extract_pdf(file_path)
        elif ext == ".docx":
            return self._extract_docx(file_path)
        elif ext == ".doc":
            return self._extract_doc(file_path)
        return ""

    def _mineru_config(self) -> dict:
        return {
            "api_key": str(conf().get("mineru_api_key", "") or os.environ.get("MINERU_TOKEN", "")).strip(),
            "api_base": str(conf().get("mineru_api_base", "") or os.environ.get("MINERU_API_BASE", "") or "https://mineru.net").strip().rstrip("/"),
            "model_version": str(conf().get("mineru_model_version", "vlm") or "vlm").strip(),
            "language": str(conf().get("mineru_language", "auto") or "auto").strip(),
            "enable_formula": bool(conf().get("mineru_enable_formula", True)),
            "enable_table": bool(conf().get("mineru_enable_table", True)),
            "enable_ocr": bool(conf().get("mineru_enable_ocr", True)),
            "timeout_seconds": int(conf().get("mineru_timeout_seconds", 1800) or 1800),
            "poll_interval_seconds": max(1, int(conf().get("mineru_poll_interval_seconds", 5) or 5)),
        }

    def _mineru_available(self) -> bool:
        return bool(self._mineru_config().get("api_key"))

    def _safe_extract_zip(self, zip_path: str, target_dir: str):
        target_dir = os.path.normpath(target_dir)
        with zipfile.ZipFile(zip_path) as zf:
            for member in zf.infolist():
                name = member.filename.replace("\\", "/")
                if not name or name.endswith("/"):
                    continue
                dest = os.path.normpath(os.path.join(target_dir, name))
                if not dest.startswith(target_dir + os.sep) and dest != target_dir:
                    logger.warning(f"[KnowledgeService] skipped unsafe MinerU zip member: {member.filename}")
                    continue
                os.makedirs(os.path.dirname(dest), exist_ok=True)
                with zf.open(member) as src, open(dest, "wb") as out:
                    shutil.copyfileobj(src, out)

    def _find_mineru_markdown(self, parsed_dir: str) -> str:
        candidates = []
        for root, _dirs, files in os.walk(parsed_dir):
            for name in files:
                if name.lower().endswith(".md"):
                    path = os.path.join(root, name)
                    score = 0
                    lower = name.lower()
                    if lower in {"full.md", "middle.md", "content.md"}:
                        score += 20
                    if os.path.getsize(path) > 0:
                        score += min(os.path.getsize(path), 1000000) / 1000000
                    candidates.append((score, path))
        if not candidates:
            return ""
        candidates.sort(key=lambda item: item[0], reverse=True)
        return candidates[0][1]

    def _pdf_page_count(self, file_path: str) -> int:
        try:
            from pypdf import PdfReader
            return len(PdfReader(file_path).pages)
        except Exception:
            pass
        try:
            import fitz
            with fitz.open(file_path) as doc:
                return len(doc)
        except Exception:
            return 0

    def _split_pdf_for_mineru(self, file_path: str, split_dir: str, max_pages: int = 200) -> list:
        os.makedirs(split_dir, exist_ok=True)
        try:
            from pypdf import PdfReader, PdfWriter
            reader = PdfReader(file_path)
            total = len(reader.pages)
            if total <= max_pages:
                return [file_path]
            parts = []
            stem = os.path.splitext(os.path.basename(file_path))[0]
            for start in range(0, total, max_pages):
                writer = PdfWriter()
                end = min(start + max_pages, total)
                for page in reader.pages[start:end]:
                    writer.add_page(page)
                part_path = os.path.join(split_dir, f"{stem}_pages_{start + 1:04d}_{end:04d}.pdf")
                with open(part_path, "wb") as f:
                    writer.write(f)
                parts.append(part_path)
            return parts
        except Exception as exc:
            logger.warning(f"[KnowledgeService] failed to split PDF for MinerU: {exc}")
            return [file_path]

    def _extract_single_pdf_with_mineru(self, file_path: str, parsed_dir: str, data_id: str) -> str:
        cfg = self._mineru_config()
        if not cfg.get("api_key"):
            return ""

        os.makedirs(parsed_dir, exist_ok=True)
        cached_md = self._find_mineru_markdown(parsed_dir)
        if cached_md:
            with open(cached_md, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
            if text.strip():
                self._emit_progress("mineru_cache", f"Using cached MinerU parse for {os.path.basename(file_path)}")
                return self._normalize_formula_blocks(text)

        try:
            import requests
        except Exception as exc:
            logger.warning(f"[KnowledgeService] MinerU unavailable because requests cannot be imported: {exc}")
            return ""

        headers = {
            "Authorization": f"Bearer {cfg['api_key']}",
            "Content-Type": "application/json",
        }
        file_name = os.path.basename(file_path)
        base = cfg["api_base"] + "/"
        create_url = urljoin(base, "api/v4/file-urls/batch")
        payload = {
            "files": [{"name": file_name, "data_id": data_id, "is_ocr": cfg["enable_ocr"]}],
            "model_version": cfg["model_version"],
            "language": cfg["language"],
            "enable_formula": cfg["enable_formula"],
            "enable_table": cfg["enable_table"],
        }
        self._emit_progress("mineru_upload_init", f"Requesting MinerU upload URL for {file_name}")
        try:
            resp = requests.post(create_url, headers=headers, json=payload, timeout=60)
            resp.raise_for_status()
            created = resp.json()
            data = created.get("data") or created
            batch_id = data.get("batch_id") or created.get("batch_id")
            upload_urls = data.get("file_urls") or data.get("upload_urls") or created.get("file_urls") or []
            upload_url = upload_urls[0] if upload_urls else data.get("upload_url")
            if isinstance(upload_url, dict):
                upload_url = upload_url.get("url") or upload_url.get("upload_url")
            if not batch_id or not upload_url:
                raise RuntimeError(f"MinerU did not return batch_id/upload_url: {str(created)[:300]}")

            self._emit_progress("mineru_uploading", f"Uploading {file_name} to MinerU", batch_id=batch_id)
            with open(file_path, "rb") as f:
                put_resp = requests.put(upload_url, data=f, timeout=300)
            put_resp.raise_for_status()

            result_url = urljoin(base, f"api/v4/extract-results/batch/{batch_id}")
            deadline = time.time() + cfg["timeout_seconds"]
            result_item = None
            while time.time() < deadline:
                self._emit_progress("mineru_processing", f"Waiting for MinerU parse: {file_name}", batch_id=batch_id)
                poll = requests.get(result_url, headers={"Authorization": f"Bearer {cfg['api_key']}"}, timeout=60)
                poll.raise_for_status()
                result_json = poll.json()
                data = result_json.get("data") or result_json
                items = data.get("extract_result") or data.get("extract_results") or data.get("results") or []
                if isinstance(items, dict):
                    items = [items]
                for item in items:
                    item_data_id = str(item.get("data_id", data_id))
                    if item_data_id and item_data_id != str(data_id):
                        continue
                    state = str(item.get("state") or item.get("status") or "").lower()
                    if state in {"done", "completed", "success"} or item.get("full_zip_url"):
                        result_item = item
                        break
                    if state in {"failed", "error"}:
                        raise RuntimeError(item.get("err_msg") or item.get("message") or "MinerU parse failed")
                if result_item:
                    break
                time.sleep(cfg["poll_interval_seconds"])
            if not result_item:
                raise TimeoutError(f"MinerU parse timed out after {cfg['timeout_seconds']} seconds")

            zip_url = result_item.get("full_zip_url") or result_item.get("zip_url") or result_item.get("download_url")
            if not zip_url:
                raise RuntimeError(f"MinerU result has no zip URL: {str(result_item)[:300]}")
            self._emit_progress("mineru_downloading", f"Downloading MinerU result for {file_name}")
            zip_resp = requests.get(zip_url, timeout=300)
            zip_resp.raise_for_status()
            zip_path = os.path.join(parsed_dir, "mineru_result.zip")
            with open(zip_path, "wb") as f:
                f.write(zip_resp.content)
            self._safe_extract_zip(zip_path, parsed_dir)
            md_path = self._find_mineru_markdown(parsed_dir)
            if not md_path:
                raise RuntimeError("MinerU result did not include a Markdown file")
            with open(md_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
            self._emit_progress("mineru_done", f"MinerU parse completed for {file_name}")
            return self._normalize_formula_blocks(text)
        except Exception as exc:
            logger.warning(f"[KnowledgeService] MinerU parse failed for {file_path}: {exc}")
            self._emit_progress("mineru_fallback", f"MinerU failed; using local PDF parser: {exc}", current_file=os.path.basename(file_path))
            return ""

    def _extract_pdf_with_mineru(self, file_path: str, book_id: str = "") -> str:
        signature = self._source_file_signature(file_path)
        source_hash = (signature.get("content_hash") or hashlib.sha256(file_path.encode("utf-8")).hexdigest())[:16]
        parsed_root = os.path.join(self._resolve_book_dir(book_id), "_parsed", "mineru", source_hash)
        os.makedirs(parsed_root, exist_ok=True)

        combined_md = os.path.join(parsed_root, "combined.md")
        if os.path.isfile(combined_md):
            with open(combined_md, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
            if text.strip():
                self._emit_progress("mineru_cache", f"Using cached MinerU parse for {os.path.basename(file_path)}")
                return self._normalize_formula_blocks(text)

        page_count = self._pdf_page_count(file_path)
        pdf_parts = [file_path]
        if page_count > 200:
            self._emit_progress(
                "mineru_split",
                f"PDF has {page_count} pages; splitting into MinerU-compatible parts",
                current_file=os.path.basename(file_path),
                total_pages=page_count,
            )
            pdf_parts = self._split_pdf_for_mineru(file_path, os.path.join(parsed_root, "splits"), max_pages=200)

        texts = []
        for idx, part_path in enumerate(pdf_parts, start=1):
            part_dir = parsed_root if len(pdf_parts) == 1 else os.path.join(parsed_root, f"part_{idx:03d}")
            part_id = source_hash if len(pdf_parts) == 1 else f"{source_hash}_part_{idx:03d}"
            self._emit_progress(
                "mineru_part",
                f"Parsing PDF part {idx}/{len(pdf_parts)} with MinerU",
                current_file=os.path.basename(part_path),
                current_part=idx,
                total_parts=len(pdf_parts),
            )
            text = self._extract_single_pdf_with_mineru(part_path, part_dir, part_id)
            if not text.strip():
                return ""
            texts.append(f"\n\n<!-- mineru_part:{idx}/{len(pdf_parts)} source:{os.path.basename(part_path)} -->\n\n{text.strip()}")

        merged = "\n\n".join(texts).strip()
        if merged:
            with open(combined_md, "w", encoding="utf-8") as f:
                f.write(merged + "\n")
        return self._normalize_formula_blocks(merged)

    def _extract_pdf(self, file_path: str, book_id: str = "") -> str:
        if self._mineru_available():
            mineru_text = self._extract_pdf_with_mineru(file_path, book_id=book_id)
            if mineru_text.strip():
                return mineru_text
        return self._extract_pdf_local(file_path)

    def _extract_pdf_local(self, file_path: str) -> str:
        try:
            import pdfplumber
            parts = []
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    tables = page.extract_tables() or []
                    if text.strip():
                        parts.append(f"\n\n<!-- page:{page_idx} -->\n{text.strip()}")
                    for table_idx, table in enumerate(tables, start=1):
                        md_table = self._table_to_markdown(table)
                        if md_table:
                            parts.append(f"\n\n<!-- page:{page_idx} table:{table_idx} -->\n{md_table}")
            if parts:
                return self._normalize_formula_blocks("\n\n".join(parts))
        except Exception:
            pass
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            parts = []
            for page_idx, page in enumerate(reader.pages, start=1):
                text = page.extract_text()
                if text:
                    parts.append(f"\n\n<!-- page:{page_idx} -->\n{text}")
            return self._normalize_formula_blocks("\n\n".join(parts))
        except Exception as e:
            logger.warning(f"[KnowledgeService] PDF extraction failed for {file_path}: {e}")
            return ""

    def _extract_docx(self, file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            parts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                md_table = self._table_to_markdown(rows)
                if md_table:
                    parts.append(md_table)
            return self._normalize_formula_blocks("\n\n".join(parts))
        except Exception as e:
            logger.warning(f"[KnowledgeService] DOCX extraction failed for {file_path}: {e}")
            return ""

    def _extract_doc(self, file_path: str) -> str:
        try:
            result = subprocess.run(
                [sys.executable, "-m", "docx2txt", "convert", file_path, "-"],
                capture_output=True, text=True, timeout=30
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout.strip()
        except Exception:
            pass
        try:
            from docx import Document
            doc = Document(file_path)
            parts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)
            return "\n\n".join(parts)
        except Exception as e:
            logger.warning(f"[KnowledgeService] DOC extraction failed for {file_path}: {e}")
        return ""

    def _table_to_markdown(self, table) -> str:
        rows = []
        for row in table or []:
            cells = [re.sub(r"\s+", " ", str(cell or "").strip()) for cell in (row or [])]
            if any(cells):
                rows.append(cells)
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        header = rows[0]
        body = rows[1:] or [[""] * width]
        lines = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(["---"] * width) + " |",
        ]
        lines.extend("| " + " | ".join(row) + " |" for row in body)
        return "\n".join(lines)

    def _normalize_formula_blocks(self, text: str) -> str:
        lines = []
        formula_re = re.compile(r"(?=.*[=+\-*/^])(?=.*[A-Za-z\u0391-\u03A9\u03B1-\u03C9])^[A-Za-z\u0391-\u03A9\u03B1-\u03C90-9\s_{}()[\].,+\-*/^=<>\u2264\u2265%:;|]+$")
        for raw in (text or "").splitlines():
            line = raw.strip()
            if line and formula_re.match(line) and len(line) <= 180 and not line.startswith("|"):
                lines.append(f"$${line}$$")
            else:
                lines.append(raw)
        return "\n".join(lines)

    def _normalization_terms(self, *values) -> list:
        """Return stable search terms for metadata recall."""
        terms = []
        seen = set()
        for value in values:
            if isinstance(value, (list, tuple, set)):
                candidates = value
            else:
                candidates = re.findall(r"[\u4e00-\u9fffA-Za-z0-9_-]{2,}", str(value or ""))
            for item in candidates:
                term = str(item or "").strip().lower()
                term = re.sub(r"[_\-\s]+", "", term)
                if len(term) < 2 or term in seen:
                    continue
                seen.add(term)
                terms.append(term)
                if len(terms) >= 80:
                    return terms
        return terms

    def _asset_filter_config(self) -> dict:
        return {
            "enabled": bool(conf().get("knowledge_extract_assets", True)),
            "skip_logo_watermark": bool(conf().get("knowledge_skip_logo_watermark_assets", True)),
            "min_width": int(conf().get("knowledge_min_asset_width", 120) or 120),
            "min_height": int(conf().get("knowledge_min_asset_height", 120) or 120),
            "min_area": int(conf().get("knowledge_min_asset_area", 20000) or 20000),
        }

    def _image_size_from_bytes(self, data: bytes) -> tuple:
        try:
            from PIL import Image
            with Image.open(BytesIO(data)) as img:
                return img.size
        except Exception:
            return (0, 0)

    def _should_keep_extracted_asset(self, data: bytes, name: str = "", width: int = 0, height: int = 0) -> tuple:
        cfg = self._asset_filter_config()
        if not cfg["enabled"]:
            return False, "asset extraction disabled"
        lower_name = (name or "").lower()
        if cfg["skip_logo_watermark"] and any(token in lower_name for token in ("logo", "watermark", "stamp", "seal", "header", "footer")):
            return False, "likely logo/watermark by name"
        if not width or not height:
            width, height = self._image_size_from_bytes(data)
        if not width or not height:
            return True, ""
        area = width * height
        if width < cfg["min_width"] or height < cfg["min_height"] or area < cfg["min_area"]:
            return False, f"small image {width}x{height}"
        aspect = max(width / max(height, 1), height / max(width, 1))
        if cfg["skip_logo_watermark"] and aspect > 8:
            return False, f"banner-like image {width}x{height}"
        return True, ""

    def _extract_assets_for_wiki(self, file_path: str, wiki_dir: str, source_id: str) -> list:
        if not self._asset_filter_config()["enabled"]:
            return []
        asset_dir = os.path.join(wiki_dir, "assets", source_id, "images")
        assets = []
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".docx":
            try:
                with zipfile.ZipFile(file_path) as zf:
                    media = [n for n in zf.namelist() if n.startswith("word/media/")]
                    for idx, name in enumerate(media, start=1):
                        suffix = os.path.splitext(name)[1].lower() or ".bin"
                        if suffix not in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".emf", ".wmf"):
                            continue
                        data = zf.read(name)
                        keep, reason = self._should_keep_extracted_asset(data, name=name)
                        if not keep:
                            logger.info(f"[KnowledgeService] DOCX image skipped ({reason}): {name}")
                            continue
                        os.makedirs(asset_dir, exist_ok=True)
                        filename = f"image_{idx:03d}{suffix}"
                        out_path = os.path.join(asset_dir, filename)
                        with open(out_path, "wb") as f:
                            f.write(data)
                        rel = os.path.relpath(out_path, wiki_dir).replace("\\", "/")
                        assets.append({
                            "id": f"{source_id}_image_{idx:03d}",
                            "path": rel,
                            "source": os.path.basename(file_path),
                            "description": "Image extracted from DOCX media folder. Review placement against the source document.",
                        })
            except Exception as e:
                logger.info(f"[KnowledgeService] DOCX image extraction skipped: {e}")
        elif ext == ".pdf":
            try:
                import fitz
                doc = fitz.open(file_path)
                image_idx = 0
                for page_idx in range(len(doc)):
                    for image_info in doc.get_page_images(page_idx):
                        xref = image_info[0]
                        extracted = doc.extract_image(xref)
                        data = extracted.get("image")
                        suffix = "." + (extracted.get("ext") or "png")
                        if not data:
                            continue
                        width = int(extracted.get("width") or image_info[2] or 0)
                        height = int(extracted.get("height") or image_info[3] or 0)
                        keep, reason = self._should_keep_extracted_asset(
                            data,
                            name=f"page_{page_idx + 1}_xref_{xref}{suffix}",
                            width=width,
                            height=height,
                        )
                        if not keep:
                            logger.info(f"[KnowledgeService] PDF image skipped ({reason}): page {page_idx + 1}, xref {xref}")
                            continue
                        image_idx += 1
                        os.makedirs(asset_dir, exist_ok=True)
                        filename = f"page_{page_idx + 1:03d}_image_{image_idx:03d}{suffix}"
                        out_path = os.path.join(asset_dir, filename)
                        with open(out_path, "wb") as f:
                            f.write(data)
                        rel = os.path.relpath(out_path, wiki_dir).replace("\\", "/")
                        assets.append({
                            "id": f"{source_id}_image_{image_idx:03d}",
                            "path": rel,
                            "source": os.path.basename(file_path),
                            "page": page_idx + 1,
                            "description": f"Image extracted from PDF page {page_idx + 1}.",
                        })
            except Exception as e:
                logger.info(f"[KnowledgeService] PDF image extraction skipped: {e}")
        return assets

    def _wiki_base_dir(self, book_id: str = "") -> str:
        return os.path.join(self._resolve_book_dir(book_id), "_llm_wiki")

    def _task_status_path(self, book_id: str = "") -> str:
        task_dir = os.path.join(self._wiki_base_dir(book_id), "tasks")
        os.makedirs(task_dir, exist_ok=True)
        return os.path.join(task_dir, "organize_status.json")

    def _load_task_status(self, book_id: str = "") -> dict:
        path = self._task_status_path(book_id)
        if os.path.isfile(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if isinstance(data, dict):
                    data.setdefault("sources", {})
                    return data
            except Exception:
                pass
        return {"version": "llm-wiki-task-v1", "book_id": book_id or "", "sources": {}}

    def _save_task_status(self, book_id: str, status: dict):
        status["version"] = "llm-wiki-task-v1"
        status["book_id"] = book_id or ""
        status["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
        path = self._task_status_path(book_id)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(status, f, ensure_ascii=False, indent=2)

    def get_organize_task_status(self, book_id: str = "", stale_after_seconds: int = 1800) -> dict:
        status = self._load_task_status(book_id)
        now = time.time()
        changed = False
        for source in (status.get("sources") or {}).values():
            if source.get("status") != "processing":
                continue
            updated_at = source.get("updated_at") or source.get("started_at") or ""
            try:
                ts = time.mktime(time.strptime(updated_at[:19], "%Y-%m-%dT%H:%M:%S"))
            except Exception:
                ts = 0
            if not ts or now - ts > stale_after_seconds:
                source["status"] = "interrupted"
                source["error"] = "Previous organize task did not finish. Re-run AI整理 to resume incremental indexing."
                source["interrupted_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
                changed = True
        if changed:
            self._save_task_status(book_id, status)
        return status

    def _update_source_task_status(self, book_id: str, source_key: str, **updates):
        status = self._load_task_status(book_id)
        source_status = status.setdefault("sources", {}).setdefault(source_key, {})
        source_status.update(updates)
        source_status["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
        self._save_task_status(book_id, status)

    def _safe_slug(self, text: str, fallback: str = "page") -> str:
        slug = re.sub(r"[^\w\u4e00-\u9fff]+", "_", (text or "").strip(), flags=re.UNICODE).strip("_").lower()
        return slug[:80] or f"{fallback}_{int(time.time())}"

    def _unique_chunk_stem(self, source_name: str, chunk_title: str, idx: int, used: set) -> str:
        source_stem = os.path.splitext(os.path.basename(source_name or "source"))[0]
        base = self._safe_slug(f"{source_stem}_{chunk_title}", f"chunk_{idx + 1:03d}")
        candidate = base
        suffix = 2
        while candidate in used:
            candidate = self._safe_slug(f"{base}_{suffix}", f"chunk_{idx + 1:03d}_{suffix}")
            suffix += 1
        used.add(candidate)
        return candidate

    def _load_wiki_index(self, book_id: str = "") -> dict:
        index_path = os.path.join(self._wiki_base_dir(book_id), "index.json")
        if os.path.isfile(index_path):
            try:
                with open(index_path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
        return {"version": "llm-wiki-v1", "sources": [], "chunks": [], "pages": [], "entities": [], "relations": []}

    def _file_content_hash(self, file_path: str) -> str:
        h = hashlib.sha256()
        with open(file_path, "rb") as f:
            for block in iter(lambda: f.read(1024 * 1024), b""):
                h.update(block)
        return h.hexdigest()

    def _source_file_signature(self, file_path: str) -> dict:
        try:
            stat = os.stat(file_path)
            return {
                "path": file_path,
                "name": os.path.basename(file_path),
                "file_size": int(stat.st_size),
                "file_mtime": int(stat.st_mtime),
                "content_hash": self._file_content_hash(file_path),
            }
        except OSError:
            return {"path": file_path, "name": os.path.basename(file_path)}

    def _is_source_already_indexed(self, file_path: str, book_id: str = "") -> bool:
        signature = self._source_file_signature(file_path)
        content_hash = signature.get("content_hash", "")
        task_status = self._load_task_status(book_id).get("sources", {})
        index = self._load_wiki_index(book_id)
        for source in index.get("sources", []):
            if content_hash and source.get("content_hash") != content_hash:
                continue
            source_key = source.get("content_hash") or source.get("hash") or source.get("id")
            source_state = task_status.get(source_key, {}).get("status", "indexed")
            if source_state == "indexed" and self._source_pipeline_complete(source, index):
                return True
        return False

    def _pipeline_versions(self) -> dict:
        return {
            "parse": str(conf().get("mineru_model_version", "local") or "local"),
            "chunk": f"h1-semantic-v2:{conf().get('knowledge_chunk_strategy', 'h1')}",
            "primary_metadata": "llm-wiki-primary-v1",
            "secondary_graph": "section-graph-v1" if bool(conf().get("knowledge_secondary_graph_enabled", True)) else "disabled",
            "graph": "llm-wiki-graph-v1",
        }

    def _find_source_record(self, index: dict, file_path: str) -> dict:
        signature = self._source_file_signature(file_path)
        content_hash = signature.get("content_hash", "")
        norm_path = os.path.normpath(file_path)
        for source in index.get("sources", []) or []:
            if content_hash and source.get("content_hash") == content_hash:
                return source
            if os.path.normpath(source.get("path", "")) == norm_path:
                return source
        return {}

    def _source_pipeline_complete(self, source: dict, index: dict) -> bool:
        if not source:
            return False
        versions = self._pipeline_versions()
        stages = source.get("pipeline_stages") or {}
        source_id = source.get("id", "")
        source_chunks = [c for c in index.get("chunks", []) or [] if c.get("source_id") == source_id]
        source_pages = [p for p in index.get("pages", []) or [] if p.get("source_id") == source_id]
        if not source_chunks or not source_pages:
            return False
        if stages.get("chunk", {}).get("version") != versions["chunk"]:
            return False
        primary_done = stages.get("primary_metadata", {}).get("version") == versions["primary_metadata"]
        if not primary_done and not source.get("primary_metadata_version") and not source_pages:
            return False
        if versions["secondary_graph"] != "disabled":
            if stages.get("secondary_graph", {}).get("version") != versions["secondary_graph"]:
                return False
        if stages.get("graph", {}).get("version") != versions["graph"]:
            return False
        return True

    def _chunk_files_exist(self, book_id: str, chunks: list) -> bool:
        wiki_dir = self._wiki_base_dir(book_id)
        for chunk in chunks or []:
            rel_path = chunk.get("path", "")
            if not rel_path or not os.path.isfile(os.path.join(wiki_dir, rel_path)):
                return False
        return True

    def _read_index_chunk_texts(self, book_id: str, chunk_records: list) -> list:
        wiki_dir = self._wiki_base_dir(book_id)
        chunks = []
        for record in chunk_records or []:
            path = os.path.join(wiki_dir, record.get("path", ""))
            if not os.path.isfile(path):
                return []
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = self._strip_wiki_frontmatter(f.read())
            heading = f"# {record.get('title', '')}".strip()
            if text.startswith(heading):
                text = text[len(heading):].strip()
            chunks.append({
                "title": record.get("title") or "Document",
                "section": record.get("section") or record.get("title") or "Document",
                "text": text,
                "part": 1,
            })
        return chunks

    @staticmethod
    def _strip_wiki_frontmatter(text: str) -> str:
        if text.startswith("---"):
            parts = text.split("---", 2)
            if len(parts) == 3:
                return parts[2].strip()
        return text.strip()

    def _mark_source_stage(self, source: dict, stage: str, version: str, **extra):
        stages = source.setdefault("pipeline_stages", {})
        stages[stage] = {
            "status": "done",
            "version": version,
            "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
            **extra,
        }

    def _save_wiki_index(self, book_id: str, index: dict):
        wiki_dir = self._wiki_base_dir(book_id)
        os.makedirs(wiki_dir, exist_ok=True)
        index.setdefault("embedding", {
            "enabled": False,
            "provider": "",
            "model": "",
            "dimension": 0,
            "index_path": "embeddings/index.json",
            "updated_at": "",
        })
        index["version"] = "llm-wiki-v1"
        index["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
        with open(os.path.join(wiki_dir, "index.json"), "w", encoding="utf-8") as f:
            json.dump(index, f, ensure_ascii=False, indent=2)
        self._save_embedding_manifest(book_id, index)
        self._save_compat_index(book_id, index)

    def _save_embedding_manifest(self, book_id: str, index: dict):
        wiki_dir = self._wiki_base_dir(book_id)
        emb_dir = os.path.join(wiki_dir, "embeddings")
        os.makedirs(emb_dir, exist_ok=True)
        manifest = {
            "version": "llm-wiki-embedding-v1",
            "enabled": bool(index.get("embedding", {}).get("enabled", False)),
            "provider": index.get("embedding", {}).get("provider", ""),
            "model": index.get("embedding", {}).get("model", ""),
            "dimension": index.get("embedding", {}).get("dimension", 0),
            "store": "reserved",
            "chunks": [
                {
                    "chunk_id": chunk.get("id", ""),
                    "source_id": chunk.get("source_id", ""),
                    "vector_id": (chunk.get("embedding") or {}).get("vector_id", ""),
                    "status": (chunk.get("embedding") or {}).get("status", "pending"),
                    "path": chunk.get("path", ""),
                }
                for chunk in index.get("chunks", [])
            ],
            "updated_at": index.get("updated_at"),
        }
        with open(os.path.join(emb_dir, "index.json"), "w", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)

    def _index_counts(self, index: dict) -> dict:
        return {
            "sources": len(index.get("sources", [])),
            "chunks": len(index.get("chunks", [])),
            "pages": len(index.get("pages", [])),
            "entities": len(index.get("entities", [])),
            "relations": len(index.get("relations", [])),
        }

    def _resolve_source_chunk_ids(self, raw_ids: list, chunk_records: list) -> list:
        linked = []
        for raw_id in raw_ids or []:
            if isinstance(raw_id, int) and 0 <= raw_id < len(chunk_records):
                linked.append(chunk_records[raw_id]["id"])
            elif isinstance(raw_id, str) and raw_id.isdigit() and int(raw_id) < len(chunk_records):
                linked.append(chunk_records[int(raw_id)]["id"])
            elif isinstance(raw_id, str) and raw_id:
                linked.append(raw_id)
        deduped = []
        for item in linked:
            if item not in deduped:
                deduped.append(item)
        return deduped

    def _prune_source_wiki_files(self, wiki_dir: str, subdir: str, source_name: str, keep_rel_paths: set) -> int:
        target_dir = os.path.join(wiki_dir, subdir)
        if not os.path.isdir(target_dir):
            return 0
        source_stem = os.path.splitext(os.path.basename(source_name or ""))[0]
        source_slug = self._safe_slug(source_stem, "source")
        removed = 0
        for name in os.listdir(target_dir):
            if not name.lower().endswith(".md"):
                continue
            rel_path = f"{subdir}/{name}".replace("\\", "/")
            if rel_path in keep_rel_paths:
                continue
            file_stem = os.path.splitext(name)[0]
            if not source_slug or not file_stem.startswith(source_slug):
                continue
            try:
                os.remove(os.path.join(target_dir, name))
                removed += 1
            except OSError as exc:
                logger.warning(f"[KnowledgeService] failed to prune stale wiki file {rel_path}: {exc}")
        return removed

    def _source_index_record(self, source: dict, book_id: str, index: dict) -> dict:
        source_id = source.get("id", "")
        source_chunks = [c for c in index.get("chunks", []) if c.get("source_id") == source_id]
        source_pages = [p for p in index.get("pages", []) if p.get("source_id") == source_id]
        source_relations = [
            r for r in index.get("relations", [])
            if any(str(cid).startswith(source_id) for cid in (r.get("source_chunk_ids") or []))
        ]
        return {
            "source_id": source_id,
            "name": source.get("name", ""),
            "path": source.get("path", ""),
            "file_size": source.get("file_size"),
            "file_mtime": source.get("file_mtime"),
            "content_hash": source.get("content_hash", ""),
            "hash": source.get("hash", ""),
            "status": "indexed",
            "canonical_index": f"{book_id}/_llm_wiki/index.json" if book_id else "_llm_wiki/index.json",
            "canonical_graph": f"{book_id}/_llm_wiki/graph.json" if book_id else "_llm_wiki/graph.json",
            "chunk_count": source.get("chunk_count", len(source_chunks)),
            "page_count": len(source_pages),
            "relation_count": len(source_relations),
            "asset_count": source.get("asset_count", len(source.get("assets", []))),
            "assets": source.get("assets", []),
            "updated_at": source.get("updated_at") or index.get("updated_at"),
        }

    def _save_compat_index(self, book_id: str, index: dict):
        """Keep a lightweight index.json at the visible knowledge root.

        LLM-WIKI stores the canonical data under ``_llm_wiki/index.json``.
        The web UI and older tools also look for ``knowledge/index.json`` or
        ``knowledge/<book_id>/index.json``, so write a compatibility index
        that points at canonical files and tracks source documents only.
        """
        target_dir = self._resolve_book_dir(book_id)
        os.makedirs(target_dir, exist_ok=True)
        counts = self._index_counts(index)
        source_records = [self._source_index_record(source, book_id, index) for source in index.get("sources", [])]
        compat = {
            "version": index.get("version", "llm-wiki-v1"),
            "updated_at": index.get("updated_at"),
            "book_id": book_id or "",
            "canonical_index": "_llm_wiki/index.json",
            "canonical_graph": "_llm_wiki/graph.json",
            "counts": counts,
            "files": source_records,
            "sources": source_records,
        }
        with open(os.path.join(target_dir, "index.json"), "w", encoding="utf-8") as f:
            json.dump(compat, f, ensure_ascii=False, indent=2)

        if book_id:
            root_index_path = os.path.join(self.knowledge_dir, "index.json")
            os.makedirs(self.knowledge_dir, exist_ok=True)
            root = {"version": "llm-wiki-v1", "updated_at": compat["updated_at"], "books": []}
            if os.path.isfile(root_index_path):
                try:
                    with open(root_index_path, "r", encoding="utf-8") as f:
                        root = json.load(f)
                except Exception:
                    pass
            books = [b for b in root.get("books", []) if b.get("book_id") != book_id]
            books.append({
                "book_id": book_id,
                "index": f"{book_id}/index.json",
                "canonical_index": f"{book_id}/_llm_wiki/index.json",
                "counts": counts,
                "files": source_records,
                "updated_at": compat["updated_at"],
            })
            root_files = []
            for book in books:
                for file_record in book.get("files", []):
                    root_files.append({"book_id": book.get("book_id", ""), **file_record})
            root.update({
                "version": "llm-wiki-v1",
                "updated_at": compat["updated_at"],
                "books": books,
                "files": root_files,
            })
            with open(root_index_path, "w", encoding="utf-8") as f:
                json.dump(root, f, ensure_ascii=False, indent=2)

    def sync_compat_index(self, book_id: str = "") -> dict:
        wiki_index_path = os.path.join(self._wiki_base_dir(book_id), "index.json")
        if os.path.isfile(wiki_index_path):
            index = self._load_wiki_index(book_id)
        else:
            index = {"version": "llm-wiki-v1", "sources": [], "chunks": [], "pages": [], "entities": [], "relations": [], "stale": True}
            index["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
        self._save_compat_index(book_id, index)
        return self._index_counts(index)

    def _chunk_for_wiki(self, content: str, max_tokens: int = None, overlap_tokens: int = None) -> list:
        strategy = str(conf().get("knowledge_chunk_strategy", "h1") or "h1").strip().lower()
        target_chars = int(conf().get("knowledge_chunk_target_chars", 6500) or 6500)
        max_chars = int(conf().get("knowledge_chunk_max_chars", 9000) or 9000)
        overlap_chars = int(conf().get("knowledge_chunk_overlap_chars", 450) or 450)
        if max_tokens is not None:
            max_chars = max(1200, int(max_tokens) * 2)
            target_chars = max(800, int(max_chars * 0.75))
        if overlap_tokens is not None:
            overlap_chars = max(0, int(overlap_tokens) * 2)
        if max_tokens is None and strategy in {"h1", "heading", "top", "chapter"}:
            blocks = self._split_markdown_top_sections(content)
            chunks = []
            for block in blocks:
                title = block.get("title") or "Document"
                text = (block.get("text") or "").strip()
                if text:
                    chunks.append({"title": title, "text": text, "section": title, "part": 1})
            return chunks

        if max_chars <= 0:
            max_chars = max(target_chars, 9000)
        blocks = self._split_markdown_sections(content)
        if not blocks:
            blocks = [{"title": "Document", "level": 1, "text": (content or "").strip()}]

        blocks = self._merge_small_wiki_sections(blocks, target_chars=target_chars, max_chars=max_chars)

        chunks = []
        for block in blocks:
            title = block.get("title") or "Document"
            text = (block.get("text") or "").strip()
            if not text:
                continue
            if len(text) <= max_chars:
                chunks.append({"title": title, "text": text, "section": title, "part": 1})
                continue
            parts = self._split_large_section(text, max_chars=max_chars, overlap_chars=overlap_chars)
            for part_idx, part in enumerate(parts, start=1):
                chunks.append({
                    "title": f"{title} ({part_idx})" if len(parts) > 1 else title,
                    "text": part,
                    "section": title,
                    "part": part_idx,
                })
        return [c for c in chunks if c["text"].strip()]

    def _split_markdown_top_sections(self, content: str) -> list:
        lines = (content or "").splitlines()
        has_semantic_top = any(
            self._is_semantic_top_title(self._heading_title(line.strip()) or line.strip())
            for line in lines
        )
        sections = []
        current_title = "Document"
        current = []
        has_top_heading = False
        top_plain_re = re.compile(
            r"^\s*((?:第\s*[\u4e00-\u9fff\d]+\s*[\u7ae0\u7bc7\u90e8]|Chapter\s+\d+)\s*[^\n]{0,100})\s*$",
            re.I,
        )
        hash_re = re.compile(r"^\s*(#{1,6})\s+(.+?)\s*$")
        for raw in lines:
            stripped = raw.strip()
            hash_match = hash_re.match(stripped)
            plain_match = top_plain_re.match(stripped)
            title = ""
            is_top = False
            if hash_match and len(hash_match.group(1)) == 1:
                title = hash_match.group(2).strip()
                is_top = bool(title) and (not has_semantic_top or self._is_semantic_top_title(title))
            elif plain_match and self._is_wiki_heading(stripped, plain_match, primary=True):
                title = plain_match.group(1).strip()
                is_top = bool(title) and self._is_semantic_top_title(title)

            if is_top:
                has_top_heading = True
                if current:
                    sections.append({"title": current_title, "level": 1, "text": "\n".join(current).strip()})
                    current = []
                current_title = title
                continue
            current.append(raw)

        if current:
            sections.append({"title": current_title, "level": 1, "text": "\n".join(current).strip()})
        if has_top_heading:
            return [s for s in sections if s.get("text", "").strip()]

        fallback = self._split_markdown_sections(content)
        if len(fallback) > 1:
            merged_text = []
            for section in fallback:
                title = section.get("title") or "Document"
                text = (section.get("text") or "").strip()
                if text:
                    merged_text.append(f"## {title}\n\n{text}" if title != "Document" else text)
            return [{"title": "Document", "level": 1, "text": "\n\n".join(merged_text).strip()}]
        return fallback or [{"title": "Document", "level": 1, "text": (content or "").strip()}]

    def _heading_title(self, line: str) -> str:
        match = re.match(r"^\s*#{1,6}\s+(.+?)\s*$", line or "")
        return match.group(1).strip() if match else ""

    def _is_semantic_top_title(self, title: str) -> bool:
        title = (title or "").strip()
        if not title:
            return False
        if re.match(r"^(第\s*[\u4e00-\u9fff\d]+\s*[\u7ae0\u7bc7\u90e8])\b", title):
            return True
        if re.match(r"^Chapter\s+\d+\b", title, re.I):
            return True
        if title in {"前言", "序言", "绪论", "引言", "导论", "附录"}:
            return True
        return False

    def _merge_small_wiki_sections(self, sections: list, target_chars: int, max_chars: int) -> list:
        merged = []
        current = None
        for section in sections or []:
            title = section.get("title") or "Document"
            text = (section.get("text") or "").strip()
            if not text:
                continue
            if len(text) >= target_chars:
                if current:
                    merged.append(current)
                    current = None
                merged.append(section)
                continue
            block_text = f"## {title}\n\n{text}" if title and title != "Document" else text
            if not current:
                current = {
                    "title": title,
                    "level": section.get("level", 1),
                    "text": block_text,
                    "section_titles": [title],
                }
                continue
            combined_len = len(current.get("text", "")) + len(block_text) + 2
            if combined_len <= max_chars:
                current["text"] = current.get("text", "").rstrip() + "\n\n" + block_text
                current.setdefault("section_titles", []).append(title)
                first = current["section_titles"][0]
                last = current["section_titles"][-1]
                current["title"] = first if first == last else f"{first} / {last}"
            else:
                merged.append(current)
                current = {
                    "title": title,
                    "level": section.get("level", 1),
                    "text": block_text,
                    "section_titles": [title],
                }
        if current:
            merged.append(current)
        return merged

    def _split_markdown_sections(self, content: str) -> list:
        heading_re = re.compile(
            r"^\s*(#{1,6})\s+(.+?)\s*$|^\s*((?:\u7b2c\s*[\u4e00-\u9fff\d]+\s*[\u7ae0\u8282\u7bc7\u90e8]|Chapter\s+\d+)\s*[^\n]{0,80})\s*$",
            re.I,
        )
        sections = []
        current_title = "Document"
        current_level = 1
        current = []
        for raw in (content or "").splitlines():
            stripped = raw.strip()
            match = heading_re.match(stripped)
            is_heading = self._is_wiki_heading(stripped, match, primary=True)
            if is_heading:
                if current:
                    sections.append({"title": current_title, "level": current_level, "text": "\n".join(current).strip()})
                    current = []
                if match.group(2):
                    current_title = match.group(2).strip()
                    current_level = len(match.group(1))
                else:
                    current_title = match.group(3).strip()
                    current_level = 2 if re.match(r"^\d+\.\d+", current_title) else 1
            else:
                current.append(raw)
        if current:
            sections.append({"title": current_title, "level": current_level, "text": "\n".join(current).strip()})
        if len(sections) <= 1:
            sections = self._split_sections_by_chinese_headings(content)
        return [s for s in sections if s.get("text", "").strip()]

    def _split_sections_by_chinese_headings(self, content: str) -> list:
        pattern = re.compile(r"(?m)^(?P<title>\s*(?:\d+(?:\.\d+){1,3})\s+[\u4e00-\u9fffA-Za-z][^\n]{2,80})$")
        matches = list(pattern.finditer(content or ""))
        matches = [m for m in matches if self._is_wiki_heading(m.group("title").strip(), m, primary=False)]
        if not matches:
            return [{"title": "Document", "level": 1, "text": (content or "").strip()}]
        sections = []
        preface = (content or "")[:matches[0].start()].strip()
        if preface:
            sections.append({"title": "Preface", "level": 1, "text": preface})
        for idx, match in enumerate(matches):
            start = match.end()
            end = matches[idx + 1].start() if idx + 1 < len(matches) else len(content or "")
            title = match.group("title").strip()
            body = (content or "")[start:end].strip()
            if body:
                sections.append({"title": title, "level": 1, "text": body})
        return sections

    def _is_wiki_heading(self, stripped: str, match, primary: bool = False) -> bool:
        if not match or not stripped or stripped.startswith("|"):
            return False
        if len(stripped) > 120:
            return False
        if stripped.startswith("#"):
            return True
        if primary:
            return bool(re.search(r"\u7b2c\s*[\u4e00-\u9fff\d]+\s*[\u7ae0\u8282\u7bc7\u90e8]|Chapter\s+\d+", stripped, re.I))
        if any(ch in stripped for ch in "=<>\u2264\u2265+-*/^"):
            return False
        chinese_count = len(re.findall(r"[\u4e00-\u9fff]", stripped))
        alpha_count = len(re.findall(r"[A-Za-z]", stripped))
        digit_count = len(re.findall(r"\d", stripped))
        if chinese_count + alpha_count < 3:
            return False
        return digit_count <= max(8, (chinese_count + alpha_count) * 2)

    def _split_large_section(self, text: str, max_chars: int, overlap_chars: int) -> list:
        blocks = self._atomic_markdown_blocks(text)
        parts = []
        current = []
        current_len = 0
        for block in blocks:
            block_len = len(block)
            if current and current_len + block_len > max_chars:
                parts.append("\n\n".join(current).strip())
                overlap = []
                overlap_len = 0
                for old in reversed(current):
                    if overlap_len + len(old) > overlap_chars:
                        break
                    overlap.insert(0, old)
                    overlap_len += len(old)
                current = overlap
                current_len = overlap_len
            if block_len > max_chars:
                if current:
                    parts.append("\n\n".join(current).strip())
                    current = []
                    current_len = 0
                start = 0
                while start < len(block):
                    end = min(start + max_chars, len(block))
                    parts.append(block[start:end].strip())
                    if end >= len(block):
                        break
                    start = max(end - overlap_chars, start + 1)
                continue
            current.append(block)
            current_len += block_len
        if current:
            parts.append("\n\n".join(current).strip())
        return [p for p in parts if p]

    def _atomic_markdown_blocks(self, text: str) -> list:
        blocks = []
        lines = (text or "").splitlines()
        idx = 0
        while idx < len(lines):
            line = lines[idx]
            stripped = line.strip()
            if stripped.startswith("|") and "|" in stripped[1:]:
                table = [line]
                idx += 1
                while idx < len(lines) and lines[idx].strip().startswith("|"):
                    table.append(lines[idx])
                    idx += 1
                blocks.append("\n".join(table))
                continue
            if stripped.startswith("$$"):
                formula = [line]
                idx += 1
                while idx < len(lines):
                    formula.append(lines[idx])
                    if lines[idx].strip().endswith("$$"):
                        idx += 1
                        break
                    idx += 1
                blocks.append("\n".join(formula))
                continue
            para = [line]
            idx += 1
            while idx < len(lines) and lines[idx].strip() and not lines[idx].strip().startswith("|") and not lines[idx].strip().startswith("$$"):
                para.append(lines[idx])
                idx += 1
            blocks.append("\n".join(para).strip())
            while idx < len(lines) and not lines[idx].strip():
                idx += 1
        return [b for b in blocks if b.strip()]

    def _chunk_skill_metadata(self, chunk: dict, source_name: str) -> dict:
        text = (chunk.get("text") or "").strip()
        title = (chunk.get("title") or os.path.splitext(source_name)[0] or "Document").strip()
        compact = re.sub(r"\s+", " ", text)
        first_sentence = re.split(r"(?<=[\u3002\uff01\uff1f.!?])\s*", compact, maxsplit=1)[0].strip()
        summary = first_sentence or compact[:180]
        if len(summary) > 220:
            summary = summary[:217].rstrip() + "..."

        words = re.findall(r"[\u4e00-\u9fffA-Za-z][\u4e00-\u9fffA-Za-z0-9_-]{1,}", compact)
        stopwords = {"the", "and", "for", "with", "that", "this", "from", "into", "\u6b63\u6587", "\u5185\u5bb9", "\u672c\u6587", "\u8fdb\u884c", "\u53ef\u4ee5", "\u4ee5\u53ca"}
        keywords = []
        for word in words:
            normalized = word.strip()
            if not normalized or normalized.lower() in stopwords:
                continue
            if normalized not in keywords:
                keywords.append(normalized)
            if len(keywords) >= 10:
                break

        topic_hint = "\u3001".join(keywords[:4]) if keywords else title
        use_when = f"Use this chunk when the task needs evidence, definitions, examples, or source details about {topic_hint}."
        if re.search(r"\u6b65\u9aa4|\u6d41\u7a0b|\u65b9\u6cd5|\u7b97\u6cd5|procedure|method|workflow", compact, re.I):
            content_type = "procedure"
        elif re.search(r"\u5b9a\u4e49|\u6982\u5ff5|\u6982\u8ff0|\u539f\u7406|definition|concept", compact, re.I):
            content_type = "concept"
        elif re.search(r"数据|实验|案例|结果|计算|表|公式|example|case|result", compact, re.I):
            content_type = "evidence"
        else:
            content_type = "reference"

        return {
            "summary": summary,
            "use_when": use_when,
            "keywords": keywords,
            "content_type": content_type,
            "source_quote": compact[:320],
        }

    def _fallback_wiki_items(self, chunks: list, source_name: str) -> dict:
        pages = []
        entities = {}
        relations = []
        for idx, chunk in enumerate(chunks):
            title = chunk["title"] if chunk["title"] != "Document" else os.path.splitext(source_name)[0]
            words = re.findall(r"[\u4e00-\u9fffA-Za-z][\u4e00-\u9fffA-Za-z0-9_-]{1,}", chunk["text"])
            keywords = []
            for word in words:
                if word not in keywords and len(word) >= 2:
                    keywords.append(word)
                if len(keywords) >= 12:
                    break
            metadata = self._chunk_skill_metadata(chunk, source_name)
            pages.append({
                "title": title,
                "summary": metadata.get("summary", chunk["text"][:240]),
                "aliases": [],
                "keywords": keywords[:6],
                "source_chunk_ids": [idx],
            })
            for kw in keywords[:6]:
                entities.setdefault(kw, {
                    "name": kw,
                    "type": "concept",
                    "description": f"Concept mentioned in {title}.",
                    "source_chunk_ids": [idx],
                })
                if kw != title:
                    relations.append({
                        "source": title,
                        "target": kw,
                        "relation": "mentions",
                        "source_chunk_ids": [idx],
                        "confidence": 0.35,
                        "evidence": chunk["text"][:160],
                    })
        return {"pages": pages, "entities": list(entities.values()), "relations": relations, "chunk_metadata": []}

    def _section_graph_windows(self, chunks: list) -> list:
        max_sections = int(conf().get("knowledge_secondary_graph_max_sections", 40) or 40)
        windows = []
        heading_re = re.compile(r"^\s*#{1,6}\s+(.+?)\s*$|^\s*(\d+(?:\.\d+){1,3}\s+[^\n]{2,100})\s*$")
        for chunk_idx, chunk in enumerate(chunks or []):
            lines = (chunk.get("text") or "").splitlines()
            current_title = chunk.get("title") or "Document"
            current = []
            for raw in lines:
                stripped = raw.strip()
                match = heading_re.match(stripped)
                title = ""
                if match:
                    title = (match.group(1) or match.group(2) or "").strip()
                    if title and self._is_semantic_top_title(title):
                        title = ""
                if title and self._is_wiki_heading(stripped, match, primary=False):
                    if current:
                        self._append_section_window(windows, chunk_idx, chunk, current_title, current)
                        if len(windows) >= max_sections:
                            return windows
                        current = []
                    current_title = title
                    continue
                current.append(raw)
            if current:
                self._append_section_window(windows, chunk_idx, chunk, current_title, current)
                if len(windows) >= max_sections:
                    return windows
        return windows

    def _append_section_window(self, windows: list, chunk_idx: int, chunk: dict, title: str, lines: list):
        text = "\n".join(lines).strip()
        if len(text) < 60:
            return
        windows.append({
            "chunk_index": chunk_idx,
            "chunk_title": chunk.get("title") or "Document",
            "section_title": title or chunk.get("title") or "Document",
            "text": text,
        })

    def _extract_secondary_graph_items(self, chunks: list, source_name: str) -> dict:
        if not bool(conf().get("knowledge_secondary_graph_enabled", True)):
            return {"entities": [], "relations": [], "pages": [], "chunk_metadata": []}
        sections = self._section_graph_windows(chunks)
        if not sections:
            return {"entities": [], "relations": [], "pages": [], "chunk_metadata": []}
        sample_chars = int(conf().get("knowledge_secondary_graph_sample_chars", 1800) or 1800)
        merged = {"entities": [], "relations": [], "pages": [], "chunk_metadata": []}
        try:
            llm = self._get_llm()
            batch_size = 8
            for start in range(0, len(sections), batch_size):
                batch = sections[start:start + batch_size]
                self._emit_progress(
                    "secondary_graph",
                    f"Extracting section graph {min(start + batch_size, len(sections))}/{len(sections)}",
                    current_chunks=min(start + batch_size, len(sections)),
                    total_chunks=len(sections),
                )
                parts = []
                for offset, section in enumerate(batch):
                    idx = start + offset
                    parts.append(
                        f"### section_index: {idx}\n"
                        f"chunk_index: {section['chunk_index']}\n"
                        f"chunk_title: {section['chunk_title']}\n"
                        f"section_title: {section['section_title']}\n\n"
                        f"{section['text'][:sample_chars]}"
                    )
                messages = [
                    {"role": "system", "content": (
                        "Extract a clean section-level knowledge graph for textbook retrieval. "
                        "Return strict JSON only with keys entities and relations. "
                        "Allowed entity types: concept, method, framework, platform, component, class_or_function, "
                        "api, metric, workflow_step, standard_or_protocol, dataset, organization, person. "
                        "Allowed relation labels: contains, belongs_to, used_for, based_on, implements, calls, depends_on, "
                        "compares_with, advantage, limitation, metric_of, evolves_to, example_of, part_of, defines. "
                        "Every entity must include name, type, description, source_chunk_ids, section_title, evidence. "
                        "Every relation must include source, target, relation, source_chunk_ids, section_title, evidence, confidence. "
                        "Use numeric chunk_index values from the input as source_chunk_ids. "
                        "Drop image hashes, file extensions, paths, Markdown debris, generic verbs, and unsupported facts."
                    )},
                    {"role": "user", "content": f"Source: {source_name}\n\nSections:\n" + "\n\n".join(parts)},
                ]
                response = llm.call(messages, temperature=0.1)
                if isinstance(response, str) and response.strip().startswith("[ERROR]"):
                    raise RuntimeError(response.strip()[:300])
                parsed = self._parse_llm_json(response)
                if isinstance(parsed, dict):
                    for key in ("entities", "relations"):
                        value = parsed.get(key, [])
                        if isinstance(value, list):
                            merged[key].extend(value)
        except Exception as exc:
            logger.warning(f"[KnowledgeService] section graph LLM extraction failed, using local candidates: {exc}")
            merged = self._fallback_section_graph_items(sections)
        if not (merged["entities"] or merged["relations"]):
            merged = self._fallback_section_graph_items(sections)
        return self._filter_graph_items(merged)

    def _fallback_section_graph_items(self, sections: list) -> dict:
        entities = []
        relations = []
        strong_terms = [
            r"\b(?:RAG|MCP|LLM|API|ReAct|Qdrant|Dify|Coze|n8n|LangChain|AutoGPT|BabyAGI|Transformer|OpenAI|Claude|GPT|Agent)\b",
            r"\b[A-Z][A-Za-z]+(?:Agent|Tool|Memory|Retriever|Planner|LLM|API)\b",
            r"(?:大语言模型|基础模型|语言模型|智能体|多智能体|智能体框架|智能体系统|上下文工程|低代码平台|向量数据库|知识库|记忆系统|检索系统|通信协议|评价体系|性能评估|深度研究智能体|智能旅行助手|赛博小镇)",
        ]
        term_re = re.compile("|".join(strong_terms), re.I)
        for section in sections:
            text = re.sub(r"\s+", " ", section.get("text", ""))
            names = []
            for match in term_re.finditer(text):
                name = match.group(0).strip(" ，。；;:：()（）[]【】")
                if name and name not in names and not self._is_noise_entity(name):
                    names.append(name)
                if len(names) >= 8:
                    break
            for name in names:
                entities.append({
                    "name": self._canonical_entity_name(name),
                    "type": self._guess_entity_type(name),
                    "description": f"{name} appears in section {section.get('section_title', '')}.",
                    "source_chunk_ids": [section.get("chunk_index")],
                    "section_title": section.get("section_title", ""),
                    "evidence": self._evidence_around(text, name),
                })
            section_title = section.get("section_title") or section.get("chunk_title")
            for name in names[:6]:
                relations.append({
                    "source": section_title,
                    "target": self._canonical_entity_name(name),
                    "relation": "contains",
                    "source_chunk_ids": [section.get("chunk_index")],
                    "section_title": section.get("section_title", ""),
                    "evidence": self._evidence_around(text, name),
                    "confidence": 0.45,
                })
        return {"entities": entities, "relations": relations, "pages": [], "chunk_metadata": []}

    def _filter_graph_items(self, extracted: dict) -> dict:
        entities = []
        for ent in extracted.get("entities", []) or []:
            name = self._canonical_entity_name(ent.get("name", ""))
            if self._is_noise_entity(name):
                continue
            ent = dict(ent)
            ent["name"] = name
            ent["type"] = self._normalize_entity_type(ent.get("type", "concept"))
            ent["evidence"] = str(ent.get("evidence") or ent.get("description") or "")[:300]
            if not ent["evidence"].strip():
                continue
            entities.append(ent)
        entity_names = {e["name"] for e in entities}
        relations = []
        for rel in extracted.get("relations", []) or []:
            source = self._canonical_entity_name(rel.get("source", ""))
            target = self._canonical_entity_name(rel.get("target", ""))
            if self._is_noise_entity(source) or self._is_noise_entity(target) or source == target:
                continue
            relation = self._normalize_relation_type(rel.get("relation", "related"))
            evidence = str(rel.get("evidence") or "")[:360]
            if not evidence.strip():
                continue
            rel = dict(rel)
            rel.update({"source": source, "target": target, "relation": relation, "evidence": evidence})
            relations.append(rel)
            for name in (source, target):
                if name not in entity_names and not self._is_noise_entity(name):
                    entities.append({
                        "name": name,
                        "type": "concept",
                        "description": f"Entity inferred from relation evidence: {evidence[:120]}",
                        "source_chunk_ids": rel.get("source_chunk_ids") or [],
                        "section_title": rel.get("section_title", ""),
                        "evidence": evidence[:300],
                    })
                    entity_names.add(name)
        return {"entities": entities, "relations": relations, "pages": extracted.get("pages", []), "chunk_metadata": extracted.get("chunk_metadata", [])}

    def _canonical_entity_name(self, name: str) -> str:
        name = re.sub(r"\s+", " ", str(name or "")).strip(" \t\r\n-_*`'\"，。；;:：()（）[]【】")
        aliases = {
            "大语言模型": "LLM",
            "大型语言模型": "LLM",
            "llm": "LLM",
            "检索增强生成": "RAG",
            "rag": "RAG",
            "模型上下文协议": "MCP",
            "mcp": "MCP",
            "人工智能智能体": "智能体",
            "agent": "智能体",
            "Agent": "智能体",
            "AI Agent": "智能体",
            "openai": "OpenAI",
            "Openai": "OpenAI",
        }
        return aliases.get(name, name)

    def _normalize_entity_type(self, value: str) -> str:
        value = str(value or "concept").strip().lower()
        if value in {"技术", "方法", "算法", "范式", "架构范式"}:
            return "method"
        if value in {"框架"}:
            return "framework"
        if value in {"平台"}:
            return "platform"
        if value in {"协议", "标准", "标准/协议"}:
            return "standard_or_protocol"
        if value in {"组件", "系统", "模块"}:
            return "component"
        if value in {"类", "函数", "类/函数", "api接口"}:
            return "class_or_function"
        if value in {"组织", "机构"}:
            return "organization"
        if value in {"人物", "作者"}:
            return "person"
        allowed = {
            "concept", "method", "framework", "platform", "component", "class_or_function",
            "api", "metric", "workflow_step", "standard_or_protocol", "dataset", "organization", "person",
        }
        mapping = {
            "技术": "method", "方法": "method", "平台": "platform", "协议": "standard_or_protocol",
            "架构范式": "method", "类": "class_or_function", "函数": "class_or_function", "系统": "component",
        }
        return mapping.get(value, value if value in allowed else "concept")

    def _normalize_relation_type(self, value: str) -> str:
        value = str(value or "related").strip().lower()
        if value in {"包含", "提到", "涉及", "核心组件", "核心属性", "组成", "组成部分", "组成关系"}:
            return "contains"
        if value in {"属于", "实例类型", "并列关系"}:
            return "belongs_to"
        if value in {"用于", "应用于", "用途"}:
            return "used_for"
        if value in {"基于", "理论基础", "依托"}:
            return "based_on"
        if value in {"实现", "实现方式"}:
            return "implements"
        if value in {"调用", "依赖关系", "依赖"}:
            return "depends_on"
        if value in {"对比", "比较"}:
            return "compares_with"
        if value in {"优点", "优势"}:
            return "advantage"
        if value in {"缺点", "局限", "限制"}:
            return "limitation"
        if value in {"评价指标", "指标"}:
            return "metric_of"
        if value in {"发展为", "演化为"}:
            return "evolves_to"
        if value in {"示例", "例子"}:
            return "example_of"
        if value in {"定义", "定义为"}:
            return "defines"
        if value in {"提出者", "发起方", "作者"}:
            return "based_on"
        mapping = {
            "包含": "contains", "属于": "belongs_to", "用于": "used_for", "基于": "based_on",
            "实现": "implements", "调用": "calls", "依赖": "depends_on", "对比": "compares_with",
            "优点": "advantage", "缺点": "limitation", "评价指标": "metric_of", "发展为": "evolves_to",
            "示例": "example_of", "组成": "part_of", "定义": "defines", "mentions": "contains",
        }
        allowed = {
            "contains", "belongs_to", "used_for", "based_on", "implements", "calls", "depends_on",
            "compares_with", "advantage", "limitation", "metric_of", "evolves_to", "example_of", "part_of", "defines",
        }
        return mapping.get(value, value if value in allowed else "contains")

    def _guess_entity_type(self, name: str) -> str:
        if re.search(r"API|class|function|函数|类", name, re.I):
            return "class_or_function"
        if re.search(r"MCP|协议", name, re.I):
            return "standard_or_protocol"
        if re.search(r"Dify|Coze|n8n|LangChain|Qdrant", name, re.I):
            return "platform"
        if re.search(r"RAG|ReAct|方法|算法|范式", name, re.I):
            return "method"
        return "concept"

    def _is_noise_entity(self, name: str) -> bool:
        name = str(name or "").strip()
        if len(name) < 2 or len(name) > 80:
            return True
        if len(re.findall(r"[\u4e00-\u9fff]", name)) > 18:
            return True
        if re.fullmatch(r"[0-9a-fA-F]{16,}", name):
            return True
        if re.fullmatch(r"\d+(?:\.\d+)*", name):
            return True
        if re.search(r"\.(?:jpg|jpeg|png|svg|gif|pdf|md|json|py|js|ts)$", name, re.I):
            return True
        if any(token in name.lower() for token in ("mineru_part", "source:", "images/", "\\", "/", "http://", "https://")):
            return True
        generic = {
            "source", "page", "result", "results", "data", "file", "image", "images", "title",
            "summary", "section", "document", "details", "text_image", "jpg", "jpeg", "png",
            "pdf", "md", "json", "the", "day", "repository", "使用", "配置", "系统", "结果", "内容", "步骤",
        }
        return name.lower() in generic or name in generic

    def _evidence_around(self, text: str, name: str, radius: int = 90) -> str:
        idx = text.lower().find(str(name).lower())
        if idx < 0:
            return text[:180]
        return text[max(0, idx - radius): idx + len(name) + radius].strip()

    def _extract_wiki_items(self, chunks: list, source_name: str) -> dict:
        if self._should_use_fast_wiki(chunks):
            self._emit_progress(
                "fast_metadata",
                f"Using fast local LLM-WIKI metadata for {len(chunks)} chunks",
                current_chunks=len(chunks),
                total_chunks=len(chunks),
            )
            return self._fallback_wiki_items(chunks, source_name)

        merged = {"pages": [], "entities": [], "relations": [], "chunk_metadata": []}
        llm_failed = False
        try:
            llm = self._get_llm()
            batch_size = 5
            for start in range(0, len(chunks), batch_size):
                batch = chunks[start:start + batch_size]
                self._emit_progress(
                    "extracting",
                    f"Extracting LLM-WIKI metadata {min(start + batch_size, len(chunks))}/{len(chunks)}",
                    current_chunks=min(start + batch_size, len(chunks)),
                    total_chunks=len(chunks),
                    current_batch=start // batch_size + 1,
                    total_batches=(len(chunks) + batch_size - 1) // batch_size,
                )
                sample_parts = []
                for offset, chunk in enumerate(batch):
                    global_idx = start + offset
                    sample_parts.append(
                        f"### chunk_index: {global_idx}\n"
                        f"title: {chunk.get('title', '')}\n"
                        f"section: {chunk.get('section', '')}\n\n"
                        f"{chunk.get('text', '')[:2600]}"
                    )
                sample = "\n\n".join(sample_parts)
                messages = [
                    {"role": "system", "content": (
                        "You build retrieval-ready LLM-WIKI metadata for textbook writing. "
                        "Return strict JSON only; no Markdown, no commentary. Keys: "
                        "chunk_metadata (chunk_index, summary, use_when, keywords, content_type, source_quote), "
                        "pages (title, summary, aliases, keywords, source_chunk_ids), "
                        "entities (name, type, description, source_chunk_ids), "
                        "relations (source, target, relation, source_chunk_ids, evidence, confidence). "
                        "Extract concrete domain entities, methods, formulas, parameters, standards, components, and results. "
                        "Keep source_chunk_ids as the numeric chunk_index values supplied above. "
                        "Do not invent facts, standards, entities, relations, page titles, or evidence not supported by the content. "
                        "If a field has no evidence, return an empty array or empty string."
                    )},
                    {"role": "user", "content": f"Source: {source_name}\n\nContent:\n{sample}"},
                ]
                response = llm.call(messages, temperature=0.2)
                if isinstance(response, str) and response.strip().startswith("[ERROR]"):
                    raise RuntimeError(response.strip()[:300])
                parsed = self._parse_llm_json(response)
                if isinstance(parsed, dict):
                    for key in merged:
                        value = parsed.get(key, [])
                        if isinstance(value, list):
                            merged[key].extend(value)
                else:
                    llm_failed = True
        except Exception as e:
            logger.warning(f"[KnowledgeService] LLM-WIKI extraction failed, using fallback: {e}")
            llm_failed = True
        if llm_failed and not (merged["pages"] or merged["entities"] or merged["relations"]):
            return self._fallback_wiki_items(chunks, source_name)
        fallback = self._fallback_wiki_items(chunks, source_name)
        if not merged["chunk_metadata"]:
            merged["chunk_metadata"] = fallback.get("chunk_metadata", [])
        if not merged["pages"]:
            merged["pages"] = fallback.get("pages", [])
        secondary = self._extract_secondary_graph_items(chunks, source_name)
        merged["entities"].extend(secondary.get("entities", []))
        merged["relations"].extend(secondary.get("relations", []))
        return self._dedupe_extracted_wiki(merged)

    def _should_use_fast_wiki(self, chunks: list) -> bool:
        mode = str(conf().get("knowledge_organize_mode", "auto")).lower()
        if mode in ("deep", "llm", "accurate"):
            return False
        if mode in ("fast", "local"):
            return True
        threshold = int(conf().get("knowledge_fast_chunk_threshold", 200) or 200)
        return len(chunks or []) > threshold

    def _dedupe_extracted_wiki(self, extracted: dict) -> dict:
        pages = {}
        for page in extracted.get("pages", []):
            title = (page.get("title") or "").strip()
            if not title:
                continue
            existing = pages.setdefault(title, {
                "title": title,
                "summary": page.get("summary", ""),
                "aliases": [],
                "keywords": [],
                "source_chunk_ids": [],
            })
            for key in ("aliases", "keywords", "source_chunk_ids"):
                for item in page.get(key) or []:
                    if item not in existing[key]:
                        existing[key].append(item)
            if len(page.get("summary", "")) > len(existing.get("summary", "")):
                existing["summary"] = page.get("summary", "")

        entities = {}
        for ent in extracted.get("entities", []):
            name = self._canonical_entity_name(ent.get("name") or "")
            if self._is_noise_entity(name):
                continue
            existing = entities.setdefault(name, {
                "name": name,
                "type": self._normalize_entity_type(ent.get("type", "concept")),
                "description": ent.get("description", ""),
                "source_chunk_ids": [],
                "sections": [],
                "evidence": "",
            })
            if len(ent.get("description", "")) > len(existing.get("description", "")):
                existing["description"] = ent.get("description", "")
            if ent.get("type") and existing.get("type") == "concept":
                existing["type"] = self._normalize_entity_type(ent.get("type"))
            section_title = ent.get("section_title")
            if section_title and section_title not in existing["sections"]:
                existing["sections"].append(section_title)
            evidence = str(ent.get("evidence", "") or "")
            if len(evidence) > len(existing.get("evidence", "")):
                existing["evidence"] = evidence[:300]
            for item in ent.get("source_chunk_ids") or []:
                if item not in existing["source_chunk_ids"]:
                    existing["source_chunk_ids"].append(item)

        relations = []
        seen_rel = set()
        for rel in extracted.get("relations", []):
            source = self._canonical_entity_name(rel.get("source") or "")
            target = self._canonical_entity_name(rel.get("target") or "")
            relation = self._normalize_relation_type(rel.get("relation") or "related")
            evidence = str(rel.get("evidence", "") or "")
            if self._is_noise_entity(source) or self._is_noise_entity(target) or source == target or not evidence.strip():
                continue
            key = (source, target, relation)
            if key in seen_rel:
                continue
            seen_rel.add(key)
            relations.append({
                "source": source,
                "target": target,
                "relation": relation,
                "source_chunk_ids": rel.get("source_chunk_ids") or [],
                "section_title": rel.get("section_title", ""),
                "evidence": evidence[:360],
                "confidence": rel.get("confidence", 0.6),
            })

        metadata = {}
        for item in extracted.get("chunk_metadata", []):
            try:
                idx = int(item.get("chunk_index"))
            except Exception:
                continue
            metadata[idx] = item
        return {
            "pages": list(pages.values()),
            "entities": list(entities.values()),
            "relations": relations,
            "chunk_metadata": list(metadata.values()),
        }

    def _write_wiki_graph(self, book_id: str, index: dict) -> dict:
        nodes = {}
        edges = []
        for page in index.get("pages", []):
            page_id = page.get("id") or page.get("path", "").replace(".md", "")
            if page_id:
                nodes[page_id] = {"id": page_id, "label": page.get("title", page_id), "category": "page"}
        for ent in index.get("entities", []):
            name = ent.get("name", "")
            if name:
                ent_id = "entity/" + self._safe_slug(name, "entity")
                nodes[ent_id] = {
                    "id": ent_id,
                    "label": name,
                    "category": ent.get("type", "entity"),
                    "description": ent.get("description", ""),
                    "sections": ent.get("sections", []),
                    "source_chunk_ids": ent.get("source_chunk_ids", []),
                    "evidence": ent.get("evidence", ""),
                }
        for chunk in index.get("chunks", []):
            chunk_id = chunk.get("id", "")
            if chunk_id:
                nodes[chunk_id] = {
                    "id": chunk_id,
                    "label": chunk.get("title") or chunk_id,
                    "category": "chunk",
                    "summary": chunk.get("summary", ""),
                }
                for ent_name in chunk.get("related_entities") or []:
                    ent_id = "entity/" + self._safe_slug(ent_name, "entity")
                    if ent_id in nodes:
                        edges.append({"source": chunk_id, "target": ent_id, "label": "mentions"})
        label_map = {}
        for node_id, node in nodes.items():
            label_map.setdefault(node.get("label", ""), node_id)
        for page in index.get("pages", []):
            page_id = page.get("id") or page.get("path", "").replace(".md", "")
            for alias in page.get("aliases") or []:
                label_map.setdefault(alias, page_id)
        for rel in index.get("relations", []):
            source = rel.get("source", "")
            target = rel.get("target", "")
            if not source or not target:
                continue
            source_id = label_map.get(source, "")
            target_id = label_map.get(target, "")
            if not source_id:
                source_id = "entity/" + self._safe_slug(source, "entity")
                nodes.setdefault(source_id, {"id": source_id, "label": source, "category": "entity"})
            if not target_id:
                target_id = "entity/" + self._safe_slug(target, "entity")
                nodes.setdefault(target_id, {"id": target_id, "label": target, "category": "entity"})
            if source_id and target_id and source_id != target_id:
                edges.append({
                    "source": source_id,
                    "target": target_id,
                    "label": rel.get("relation", "related"),
                    "source_chunk_ids": rel.get("source_chunk_ids") or [],
                    "section_title": rel.get("section_title", ""),
                    "evidence": rel.get("evidence", ""),
                    "confidence": rel.get("confidence", 0.6),
                })
        graph = {"nodes": list(nodes.values()), "edges": edges}
        wiki_dir = self._wiki_base_dir(book_id)
        os.makedirs(wiki_dir, exist_ok=True)
        with open(os.path.join(wiki_dir, "graph.json"), "w", encoding="utf-8") as f:
            json.dump(graph, f, ensure_ascii=False, indent=2)
        return graph

    def migrate_graph_normalization(self, book_id: str = "") -> dict:
        index = self._load_wiki_index(book_id)
        before = {
            "entities": len(index.get("entities", []) or []),
            "relations": len(index.get("relations", []) or []),
        }

        entities = {}
        for ent in index.get("entities", []) or []:
            name = self._canonical_entity_name(ent.get("name", ""))
            if self._is_noise_entity(name):
                continue
            normalized = dict(ent)
            normalized["name"] = name
            normalized["type"] = self._normalize_entity_type(ent.get("type", "concept"))
            normalized["source_chunk_ids"] = self._dedupe_list(ent.get("source_chunk_ids") or [])
            normalized["sections"] = self._dedupe_list((ent.get("sections") or []) + ([ent.get("section_title")] if ent.get("section_title") else []))
            normalized["evidence"] = str(ent.get("evidence") or "")[:300]
            self._merge_entity_record(entities, normalized)

        relation_map = {}
        for rel in index.get("relations", []) or []:
            source = self._canonical_entity_name(rel.get("source", ""))
            target = self._canonical_entity_name(rel.get("target", ""))
            if self._is_noise_entity(source) or self._is_noise_entity(target) or source == target:
                continue
            relation = self._normalize_relation_type(rel.get("relation", "related"))
            evidence = str(rel.get("evidence") or "")[:360]
            section_title = rel.get("section_title", "") or ""
            key = (source, target, relation, section_title)
            existing = relation_map.setdefault(key, {
                "source": source,
                "target": target,
                "relation": relation,
                "source_chunk_ids": [],
                "section_title": section_title,
                "evidence": "",
                "confidence": 0.0,
            })
            for cid in rel.get("source_chunk_ids") or []:
                if cid not in existing["source_chunk_ids"]:
                    existing["source_chunk_ids"].append(cid)
            if len(evidence) > len(existing.get("evidence", "")):
                existing["evidence"] = evidence
            try:
                existing["confidence"] = max(float(existing.get("confidence", 0.0)), float(rel.get("confidence", 0.6)))
            except Exception:
                existing["confidence"] = existing.get("confidence", 0.6)
            for endpoint in (source, target):
                if endpoint not in entities and not self._is_noise_entity(endpoint):
                    self._merge_entity_record(entities, {
                        "name": endpoint,
                        "type": "concept",
                        "description": "",
                        "source_chunk_ids": existing["source_chunk_ids"],
                        "sections": [section_title] if section_title else [],
                        "evidence": evidence[:300],
                    })

        index["entities"] = list(entities.values())
        index["relations"] = list(relation_map.values())
        index["graph_migration"] = {
            "version": "graph-normalize-v1",
            "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
            "before": before,
            "after": {
                "entities": len(index["entities"]),
                "relations": len(index["relations"]),
            },
        }
        self._save_wiki_index(book_id, index)
        graph = self._write_wiki_graph(book_id, index)
        return {
            "status": "success",
            "before": before,
            "after": index["graph_migration"]["after"],
            "graph_nodes": len(graph.get("nodes", [])),
            "graph_edges": len(graph.get("edges", [])),
        }

    def _merge_entity_record(self, entities: dict, ent: dict):
        name = ent.get("name", "")
        if not name:
            return
        existing = entities.setdefault(name, {
            "name": name,
            "type": ent.get("type", "concept"),
            "description": ent.get("description", ""),
            "source_chunk_ids": [],
            "sections": [],
            "evidence": "",
        })
        if ent.get("type") and existing.get("type") == "concept":
            existing["type"] = ent.get("type")
        if len(ent.get("description", "") or "") > len(existing.get("description", "") or ""):
            existing["description"] = ent.get("description", "")
        if len(ent.get("evidence", "") or "") > len(existing.get("evidence", "") or ""):
            existing["evidence"] = ent.get("evidence", "")[:300]
        if ent.get("source_id") and not existing.get("source_id"):
            existing["source_id"] = ent.get("source_id")
        existing["source_chunk_ids"] = self._dedupe_list(existing.get("source_chunk_ids", []) + (ent.get("source_chunk_ids") or []))
        existing["sections"] = self._dedupe_list(existing.get("sections", []) + (ent.get("sections") or []))

    @staticmethod
    def _dedupe_list(items: list) -> list:
        deduped = []
        for item in items or []:
            if item and item not in deduped:
                deduped.append(item)
        return deduped

    def _resume_indexed_source(self, file_path: str, book_id: str = "") -> Optional[dict]:
        index = self._load_wiki_index(book_id)
        source = self._find_source_record(index, file_path)
        if not source:
            return None
        source_id = source.get("id", "")
        chunk_records = [c for c in index.get("chunks", []) or [] if c.get("source_id") == source_id]
        page_records = [p for p in index.get("pages", []) or [] if p.get("source_id") == source_id]
        if not chunk_records or not self._chunk_files_exist(book_id, chunk_records):
            return None

        versions = self._pipeline_versions()
        resumed = False
        self._mark_source_stage(source, "parse", versions["parse"])
        if source.get("pipeline_stages", {}).get("chunk", {}).get("version") != versions["chunk"]:
            self._mark_source_stage(source, "chunk", versions["chunk"], chunk_count=len(chunk_records))
            resumed = True
        if page_records and source.get("pipeline_stages", {}).get("primary_metadata", {}).get("version") != versions["primary_metadata"]:
            self._mark_source_stage(source, "primary_metadata", versions["primary_metadata"], page_count=len(page_records))
            resumed = True

        secondary_needed = (
            versions["secondary_graph"] != "disabled"
            and source.get("pipeline_stages", {}).get("secondary_graph", {}).get("version") != versions["secondary_graph"]
        )
        if secondary_needed:
            chunks = self._read_index_chunk_texts(book_id, chunk_records)
            if not chunks:
                return None
            self._emit_progress(
                "resume_secondary_graph",
                f"Resuming section-level graph for {os.path.basename(file_path)}",
                current_file=os.path.basename(file_path),
                total_chunks=len(chunks),
            )
            secondary = self._dedupe_extracted_wiki(
                self._extract_secondary_graph_items(chunks, os.path.basename(file_path))
            )
            entity_map = {e.get("name"): e for e in index.get("entities", []) if e.get("name")}
            for ent in secondary.get("entities", []):
                if not ent.get("name"):
                    continue
                ent["source_id"] = source_id
                ent["source_chunk_ids"] = self._resolve_source_chunk_ids(ent.get("source_chunk_ids") or [], chunk_records)
                ent["sections"] = ent.get("sections") or ([ent.get("section_title")] if ent.get("section_title") else [])
                existing = entity_map.get(ent["name"])
                if existing:
                    existing_sections = existing.setdefault("sections", [])
                    for section in ent.get("sections") or []:
                        if section and section not in existing_sections:
                            existing_sections.append(section)
                    existing_ids = existing.setdefault("source_chunk_ids", [])
                    for cid in ent.get("source_chunk_ids") or []:
                        if cid not in existing_ids:
                            existing_ids.append(cid)
                    if len(ent.get("description", "")) > len(existing.get("description", "")):
                        existing["description"] = ent.get("description", "")
                    if len(ent.get("evidence", "")) > len(existing.get("evidence", "")):
                        existing["evidence"] = ent.get("evidence", "")
                else:
                    entity_map[ent["name"]] = ent
            index["entities"] = list(entity_map.values())

            rel_keys = {(r.get("source"), r.get("target"), r.get("relation"), r.get("section_title", "")) for r in index.get("relations", [])}
            for rel in secondary.get("relations", []):
                rel["source_id"] = source_id
                rel["source_chunk_ids"] = self._resolve_source_chunk_ids(rel.get("source_chunk_ids") or [], chunk_records)
                key = (rel.get("source"), rel.get("target"), rel.get("relation"), rel.get("section_title", ""))
                if key not in rel_keys and key[0] and key[1]:
                    index.setdefault("relations", []).append(rel)
                    rel_keys.add(key)
            self._mark_source_stage(
                source,
                "secondary_graph",
                versions["secondary_graph"],
                enabled=True,
                entity_count=len(secondary.get("entities", [])),
                relation_count=len(secondary.get("relations", [])),
            )
            resumed = True
        elif versions["secondary_graph"] == "disabled":
            self._mark_source_stage(source, "secondary_graph", versions["secondary_graph"], enabled=False)

        self._mark_source_stage(source, "graph", versions["graph"])
        source["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
        if self._source_pipeline_complete(source, index):
            self._save_wiki_index(book_id, index)
            self._write_wiki_graph(book_id, index)
            return {
                "entries": [],
                "organized_count": 0,
                "chunks": len(chunk_records),
                "entities": len(index.get("entities", [])),
                "relations": len(index.get("relations", [])),
                "resumed": resumed,
                "skipped_existing": not resumed,
                "stage": "resume" if resumed else "indexed",
            }
        return None

    def _build_llm_wiki(self, file_path: str, content: str, book_id: str = "") -> dict:
        wiki_dir = self._wiki_base_dir(book_id)
        os.makedirs(os.path.join(wiki_dir, "chunks"), exist_ok=True)
        os.makedirs(os.path.join(wiki_dir, "pages"), exist_ok=True)

        source_name = os.path.basename(file_path)
        self._emit_progress("chunking", f"Chunking {source_name}", current_file=source_name)
        signature = self._source_file_signature(file_path)
        content_hash = signature.get("content_hash") or hashlib.sha256(content.encode("utf-8", errors="ignore")).hexdigest()
        signature.setdefault("content_hash", content_hash)
        signature.setdefault("name", source_name)
        signature.setdefault("path", file_path)
        source_hash = content_hash[:12]
        source_id = f"{self._safe_slug(source_name, 'source')}_{source_hash}"
        chunks = self._chunk_for_wiki(content)
        fast_mode = self._should_use_fast_wiki(chunks)
        assets = self._extract_assets_for_wiki(file_path, wiki_dir, source_id)
        if fast_mode:
            self._emit_progress(
                "fast_metadata",
                f"Large source detected ({len(chunks)} chunks). Building fast local index first.",
                current_chunks=len(chunks),
                total_chunks=len(chunks),
            )
        extracted = self._extract_wiki_items(chunks, source_name)
        metadata_by_idx = {}
        for item in extracted.get("chunk_metadata", []):
            try:
                metadata_by_idx[int(item.get("chunk_index"))] = item
            except Exception:
                continue

        chunk_records = []
        used_chunk_stems = set()
        for idx, chunk in enumerate(chunks):
            if idx % 10 == 0:
                self._emit_progress("writing_chunks", f"Writing chunks {idx + 1}/{len(chunks)}", current_chunks=idx + 1, total_chunks=len(chunks))
            chunk_stem = self._unique_chunk_stem(source_name, chunk.get("title", ""), idx, used_chunk_stems)
            chunk_id = f"{source_id}_{idx + 1:03d}"
            chunk_rel = f"chunks/{chunk_stem}.md"
            metadata = self._chunk_skill_metadata(chunk, source_name)
            llm_metadata = metadata_by_idx.get(idx) or {}
            for key in ("summary", "use_when", "keywords", "content_type", "source_quote"):
                if llm_metadata.get(key):
                    metadata[key] = llm_metadata[key]
            related_entities = []
            for ent in extracted.get("entities", []):
                raw_ids = ent.get("source_chunk_ids") or []
                if idx in raw_ids or str(idx) in [str(x) for x in raw_ids]:
                    related_entities.append(ent.get("name"))
            related_entities = [e for e in related_entities if e]
            chunk_assets = [
                asset for asset in assets
                if asset.get("page") and f"page:{asset.get('page')}" in chunk.get("text", "")
            ]
            with open(os.path.join(wiki_dir, chunk_rel), "w", encoding="utf-8") as f:
                f.write(
                    "---\n"
                    f"id: {chunk_id}\n"
                    f"source: {source_name}\n"
                    f"title: {chunk['title']}\n"
                    f"section: {chunk.get('section', chunk['title'])}\n"
                    f"summary: {metadata['summary']}\n"
                    f"use_when: {metadata['use_when']}\n"
                    f"content_type: {metadata['content_type']}\n"
                    "keywords:\n"
                    + "".join(f"  - {kw}\n" for kw in metadata["keywords"])
                    + ("related_entities:\n" + "".join(f"  - {ent}\n" for ent in related_entities) if related_entities else "")
                    + ("assets:\n" + "".join(f"  - {asset['path']}\n" for asset in chunk_assets) if chunk_assets else "")
                    + "---\n\n"
                    f"# {chunk['title']}\n\n{chunk['text']}"
                )
            chunk_records.append({
                "id": chunk_id,
                "source_id": source_id,
                "title": chunk["title"],
                "section": chunk.get("section", chunk["title"]),
                "path": chunk_rel,
                "chars": len(chunk["text"]),
                "summary": metadata["summary"],
                "use_when": metadata["use_when"],
                "keywords": metadata["keywords"],
                "normalized_terms": self._normalization_terms(
                    chunk["title"],
                    chunk.get("section", chunk["title"]),
                    metadata["summary"],
                    metadata["use_when"],
                    metadata["keywords"],
                    related_entities,
                    metadata.get("source_quote", ""),
                ),
                "content_type": metadata["content_type"],
                "source_quote": metadata.get("source_quote", ""),
                "related_entities": related_entities,
                "assets": chunk_assets,
                "embedding": {
                    "status": "pending",
                    "provider": "",
                    "model": "",
                    "vector_id": "",
                    "updated_at": "",
                },
            })

        pruned_chunks = self._prune_source_wiki_files(
            wiki_dir,
            "chunks",
            source_name,
            {record.get("path", "") for record in chunk_records},
        )
        if pruned_chunks:
            self._emit_progress("pruning_chunks", f"Pruned {pruned_chunks} stale chunk files", current_file=source_name)

        index = self._load_wiki_index(book_id)
        same_source_ids = {
            s.get("id") for s in index.get("sources", [])
            if s.get("id") and (
                s.get("id") == source_id
                or (content_hash and s.get("content_hash") == content_hash)
                or os.path.normpath(s.get("path", "")) == os.path.normpath(file_path)
            )
        }
        same_source_ids.add(source_id)
        source_stem = os.path.splitext(source_name)[0]

        def belongs_to_current_source(item: dict) -> bool:
            if item.get("source_id") in same_source_ids:
                return True
            raw_ids = item.get("source_chunk_ids") or []
            if not item.get("source_id") and raw_ids and all(str(cid).isdigit() for cid in raw_ids):
                return True
            for cid in item.get("source_chunk_ids") or []:
                if str(cid).startswith(tuple(same_source_ids)):
                    return True
            text = " ".join(
                str(item.get(key, ""))
                for key in ("source", "target", "description", "evidence")
            )
            return bool(source_name and source_name in text) or bool(source_stem and source_stem in text)

        index["sources"] = [s for s in index.get("sources", []) if s.get("id") not in same_source_ids]
        index["chunks"] = [c for c in index.get("chunks", []) if c.get("source_id") not in same_source_ids]
        index["pages"] = [p for p in index.get("pages", []) if p.get("source_id") not in same_source_ids]
        index["entities"] = [e for e in index.get("entities", []) if not belongs_to_current_source(e)]
        index["relations"] = [r for r in index.get("relations", []) if not belongs_to_current_source(r)]
        index["sources"].append({
            "id": source_id,
            "name": source_name,
            "path": file_path,
            "hash": source_hash,
            **signature,
            "chunk_count": len(chunk_records),
            "asset_count": len(assets),
            "assets": assets,
            "pipeline_stages": {
                "parse": {
                    "status": "done",
                    "version": self._pipeline_versions()["parse"],
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                },
                "chunk": {
                    "status": "done",
                    "version": self._pipeline_versions()["chunk"],
                    "chunk_count": len(chunk_records),
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                },
                "primary_metadata": {
                    "status": "done",
                    "version": self._pipeline_versions()["primary_metadata"],
                    "page_count": len(extracted.get("pages", [])),
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                },
                "secondary_graph": {
                    "status": "done",
                    "version": self._pipeline_versions()["secondary_graph"],
                    "enabled": bool(conf().get("knowledge_secondary_graph_enabled", True)),
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                },
                "graph": {
                    "status": "done",
                    "version": self._pipeline_versions()["graph"],
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                },
            },
            "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        })
        index["chunks"].extend(chunk_records)

        page_records = []
        for page in extracted.get("pages", []):
            title = page.get("title") or source_name
            slug = self._safe_slug(title, "page")
            page_id = "page/" + slug
            page_rel = f"pages/{slug}.md"
            linked_chunks = self._resolve_source_chunk_ids(page.get("source_chunk_ids") or [], chunk_records)
            md = f"# {title}\n\n{page.get('summary', '')}\n\n"
            keywords = page.get("keywords") or []
            if keywords:
                md += "## Keywords\n" + "\n".join(f"- {kw}" for kw in keywords) + "\n\n"
            if linked_chunks:
                md += "## Sources\n" + "\n".join(f"- {cid}" for cid in linked_chunks) + "\n"
            with open(os.path.join(wiki_dir, page_rel), "w", encoding="utf-8") as f:
                f.write(md)
            page_records.append({
                "id": page_id,
                "title": title,
                "path": page_rel,
                "summary": page.get("summary", ""),
                "aliases": page.get("aliases") or [],
                "keywords": keywords,
                "source_id": source_id,
                "source_chunk_ids": linked_chunks,
            })

        index["pages"] = [p for p in index.get("pages", []) if p.get("source_id") not in same_source_ids] + page_records
        entity_map = {e.get("name"): e for e in index.get("entities", []) if e.get("name")}
        for ent in extracted.get("entities", []):
            if ent.get("name"):
                ent["source_id"] = source_id
                ent["source_chunk_ids"] = self._resolve_source_chunk_ids(ent.get("source_chunk_ids") or [], chunk_records)
                ent["sections"] = ent.get("sections") or ([ent.get("section_title")] if ent.get("section_title") else [])
                entity_map[ent["name"]] = ent
        index["entities"] = list(entity_map.values())
        rel_keys = {(r.get("source"), r.get("target"), r.get("relation")) for r in index.get("relations", [])}
        for rel in extracted.get("relations", []):
            key = (rel.get("source"), rel.get("target"), rel.get("relation"))
            if key not in rel_keys and key[0] and key[1]:
                rel["source_id"] = source_id
                rel["source_chunk_ids"] = self._resolve_source_chunk_ids(rel.get("source_chunk_ids") or [], chunk_records)
                index.setdefault("relations", []).append(rel)
                rel_keys.add(key)
        index["embedding"] = index.get("embedding") or {
            "enabled": False,
            "provider": "",
            "model": "",
            "dimension": 0,
            "index_path": "embeddings/index.json",
            "updated_at": "",
        }
        self._save_wiki_index(book_id, index)
        self._emit_progress("graph", "Building knowledge graph", current_file=source_name)
        self._write_wiki_graph(book_id, index)
        return {
            "entries": page_records,
            "organized_count": len(page_records),
            "chunks": len(chunk_records),
            "entities": len(extracted.get("entities", [])),
            "relations": len(extracted.get("relations", [])),
            "assets": len(assets),
            "wiki_dir": wiki_dir,
        }

    def parse_document(self, file_path: str, book_id: str = "") -> dict:
        if not file_path or ".." in file_path:
            raise ValueError("invalid file path")
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"file not found: {file_path}")
        resumed = self._resume_indexed_source(file_path, book_id)
        if resumed:
            return resumed
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".md":
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
        elif ext in (".txt", ".csv", ".json"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif ext == ".pdf":
            content = self._extract_pdf(file_path, book_id=book_id)
            if not content.strip():
                content = f"[binary document: {os.path.basename(file_path)}]"
        elif ext in (".doc", ".docx"):
            content = self._extract_text(file_path)
            if not content.strip():
                content = f"[binary document: {os.path.basename(file_path)}]"
        else:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        if not content.strip():
            return {"entries": [], "organized_count": 0, "chunks": 0}
        return self._build_llm_wiki(file_path, content, book_id)

    def _get_llm(self):
        from bridge.textbook_bridge import _LightweightLLM
        return _LightweightLLM(role="knowledge")

    def _parse_llm_json(self, response: str) -> list:
        text = response.strip()
        if text.startswith("```"):
            lines = text.split("\n")
            text = "\n".join(lines[1:-1] if lines[-1].strip() == "```" else lines[1:])
        try:
            result = json.loads(text)
            if isinstance(result, list):
                return result
            if isinstance(result, dict) and any(k in result for k in ("pages", "entities", "relations")):
                return result
            if isinstance(result, dict) and "entries" in result:
                return result["entries"]
        except json.JSONDecodeError:
            pass
        json_match = re.search(r'\[[\s\S]*\]', text)
        if json_match:
            try:
                return json.loads(json_match.group())
            except json.JSONDecodeError:
                pass
        return []

    def _fallback_parse(self, content: str) -> list:
        entries = []
        lines = content.split("\n")
        current_title = ""
        current_lines = []
        for line in lines:
            stripped = line.strip()
            if stripped.startswith("# ") or stripped.startswith("## "):
                if current_title and current_lines:
                    body = "\n".join(current_lines).strip()
                    if body:
                        entries.append({
                            "title": current_title,
                            "summary": body[:100],
                            "keywords": [],
                            "related_concepts": [],
                            "category": "concepts"
                        })
                current_title = stripped.lstrip("#").strip()
                current_lines = []
            else:
                if stripped:
                    current_lines.append(stripped)
        if current_title and current_lines:
            body = "\n".join(current_lines).strip()
            if body:
                entries.append({
                    "title": current_title,
                    "summary": body[:100],
                    "keywords": [],
                    "related_concepts": [],
                    "category": "concepts"
                })
        return entries

    def _save_entries(self, entries: list, book_id: str = "") -> list:
        saved = []
        if book_id:
            base_dir = os.path.join(self.knowledge_dir, book_id)
        else:
            base_dir = self.knowledge_dir
        for entry in entries:
            title = entry.get("title", "").strip()
            if not title:
                continue
            category = entry.get("category", "concepts").strip().lower()
            category = re.sub(r'[^\w]', '_', category) if category else "concepts"
            slug = re.sub(r'[^\w\u4e00-\u9fff]+', '_', title).strip('_').lower()
            if not slug:
                slug = f"entry_{int(time.time())}"
            target_dir = os.path.join(base_dir, category)
            os.makedirs(target_dir, exist_ok=True)
            file_path = os.path.join(target_dir, f"{slug}.md")
            keywords = entry.get("keywords", [])
            related = entry.get("related_concepts", [])
            summary = entry.get("summary", "")
            md_content = f"# {title}\n\n"
            if summary:
                md_content += f"{summary}\n\n"
            if keywords:
                md_content += "**关键词**: " + ", ".join(keywords) + "\n\n"
            if related:
                md_content += "**关联概念**: " + ", ".join(related) + "\n\n"
            try:
                with open(file_path, "w", encoding="utf-8") as f:
                    f.write(md_content)
                saved.append(entry)
            except Exception as e:
                logger.warning(f"[KnowledgeService] failed to save entry '{title}': {e}")
        return saved

    def organize_knowledge(self, book_id: str = "", force: bool = False) -> dict:
        base = self._resolve_book_dir(book_id)
        if not os.path.isdir(base):
            return {"status": "success", "message": "knowledge dir not found", "organized_count": 0}
        self._emit_progress("scanning", "Scanning knowledge sources", force=force)
        organized_categories = {"concepts", "methods", "entities", "principles", "standards", "facts", "procedures"}
        all_files = []
        for root, dirs, files in os.walk(base):
            dirs[:] = [d for d in dirs if not d.startswith(".") and d not in {"_llm_wiki", "_parsed", "_raw"}]
            rel = os.path.relpath(root, base).replace("\\", "/")
            if rel != ".":
                first_part = rel.split("/")[0]
                if first_part in organized_categories:
                    dirs[:] = []
                    continue
            for fname in files:
                if fname.startswith(".") or fname in {"index.md", "index.json", "log.md", "cross_references.md"}:
                    continue
                full_path = os.path.join(root, fname)
                all_files.append(full_path)
        pending_files = []
        skipped_count = 0
        task_status = self._load_task_status(book_id)
        for fp in all_files:
            ext = os.path.splitext(fp)[1].lower()
            if ext not in (".md", ".txt", ".csv", ".json", ".pdf", ".doc", ".docx"):
                continue
            if not force and self._is_source_already_indexed(fp, book_id):
                skipped_count += 1
                continue
            pending_files.append(fp)
        scan_message = (
            f"Found {len(pending_files)} files to reorganize, skipped {skipped_count} indexed files"
            if force
            else f"Found {len(pending_files)} new or changed files, skipped {skipped_count} indexed files"
        )
        self._emit_progress(
            "scanned",
            scan_message,
            total_files=len(pending_files),
            skipped_files=skipped_count,
            force=force,
        )
        organized_count = 0
        processed_count = 0
        for fp in pending_files:
            processed_count += 1
            signature = self._source_file_signature(fp)
            source_key = signature.get("content_hash") or os.path.normpath(fp)
            self._update_source_task_status(
                book_id,
                source_key,
                status="processing",
                name=os.path.basename(fp),
                path=fp,
                file_size=signature.get("file_size"),
                file_mtime=signature.get("file_mtime"),
                content_hash=signature.get("content_hash", ""),
                started_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
                error="",
            )
            self._emit_progress(
                "processing_file",
                f"Processing {os.path.basename(fp)} ({processed_count}/{len(pending_files)})",
                current_file=os.path.basename(fp),
                current_file_index=processed_count,
                total_files=len(pending_files),
            )
            try:
                result = self.parse_document(fp, book_id)
                organized_count += result.get("organized_count", 0)
                self._update_source_task_status(
                    book_id,
                    source_key,
                    status="indexed",
                    organized_count=result.get("organized_count", 0),
                    chunk_count=result.get("chunks", 0),
                    entity_count=result.get("entities", 0),
                    relation_count=result.get("relations", 0),
                    finished_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
                )
            except Exception as e:
                logger.warning(f"[KnowledgeService] organize failed for {fp}: {e}")
                self._update_source_task_status(
                    book_id,
                    source_key,
                    status="failed",
                    error=str(e),
                    finished_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
                )
                self._emit_progress("file_error", str(e), current_file=os.path.basename(fp))
        self._emit_progress("cross_references", "Writing cross references")
        graph_result = self.build_knowledge_graph(book_id)
        cross_ref_count = self._write_cross_references(graph_result, book_id)
        self._emit_progress("done", "Knowledge organization complete", total_files=len(pending_files), skipped_files=skipped_count, force=force)
        return {
            "status": "success",
            "message": f"organized {organized_count} entries, skipped {skipped_count} indexed files, {cross_ref_count} cross-references",
            "organized_count": organized_count,
            "processed_files": len(pending_files),
            "skipped_files": skipped_count,
            "cross_references": cross_ref_count,
            "force": force,
        }

    def _write_cross_references(self, graph_data: dict, book_id: str = "") -> int:
        edges = graph_data.get("edges", [])
        if not edges:
            return 0
        if book_id:
            base_dir = os.path.join(self.knowledge_dir, book_id)
        else:
            base_dir = self.knowledge_dir
        cross_ref_path = os.path.join(base_dir, "cross_references.md")
        content = "# 知识关联索引\n\n"
        for edge in edges:
            source = edge.get("source", "")
            target = edge.get("target", "")
            label = edge.get("label", "")
            content += f"- [{source}]({source}) → [{target}]({target})"
            if label:
                content += f" ({label})"
            content += "\n"
        try:
            with open(cross_ref_path, "w", encoding="utf-8") as f:
                f.write(content)
        except Exception as e:
            logger.warning(f"[KnowledgeService] failed to write cross references: {e}")
        return len(edges)

    def build_knowledge_graph(
        self,
        book_id: str = "",
        limit: int = 140,
        focus_id: str = "",
        depth: int = 1,
        query: str = "",
        min_confidence: float = 0.0,
    ) -> dict:
        base = self._resolve_book_dir(book_id)
        knowledge_path = Path(base)
        if not knowledge_path.is_dir():
            return {"nodes": [], "edges": []}
        wiki_graph = os.path.join(self._wiki_base_dir(book_id), "graph.json")
        if os.path.isfile(wiki_graph):
            try:
                with open(wiki_graph, "r", encoding="utf-8") as f:
                    graph = json.load(f)
                    return self._project_graph(
                        graph.get("nodes", []),
                        graph.get("edges") or graph.get("links") or [],
                        limit=limit,
                        focus_id=focus_id,
                        depth=depth,
                        query=query,
                        min_confidence=min_confidence,
                    )
            except Exception:
                pass
        nodes = []
        edges = []
        keyword_map = {}
        concept_map = {}
        for md_file in knowledge_path.rglob("*.md"):
            rel = str(md_file.relative_to(knowledge_path)).replace("\\", "/")
            if rel in ("index.md", "log.md", "cross_references.md"):
                continue
            parts = rel.split("/")
            category = parts[0] if len(parts) > 1 else "root"
            title = md_file.stem.replace("-", " ").title()
            keywords = []
            related = []
            try:
                content = md_file.read_text(encoding="utf-8")
                first_line = content.strip().split("\n")[0]
                if first_line.startswith("# "):
                    title = first_line[2:].strip()
                kw_match = re.search(r'\*\*关键词\*\*:\s*(.+)', content)
                if kw_match:
                    keywords = [k.strip() for k in kw_match.group(1).split(",") if k.strip()]
                rel_match = re.search(r'\*\*关联概念\*\*:\s*(.+)', content)
                if rel_match:
                    related = [r.strip() for r in rel_match.group(1).split(",") if r.strip()]
            except Exception:
                pass
            node_id = rel.replace(".md", "")
            nodes.append({"id": node_id, "label": title, "category": category})
            for kw in keywords:
                keyword_map.setdefault(kw, []).append(node_id)
            for concept in related:
                concept_map.setdefault(concept, []).append(node_id)
        node_ids = {n["id"] for n in nodes}
        seen_edges = set()
        for kw, node_list in keyword_map.items():
            for i in range(len(node_list)):
                for j in range(i + 1, len(node_list)):
                    if node_list[i] in node_ids and node_list[j] in node_ids:
                        key = tuple(sorted([node_list[i], node_list[j]]))
                        if key not in seen_edges:
                            seen_edges.add(key)
                            edges.append({"source": node_list[i], "target": node_list[j], "label": kw})
        for concept, node_list in concept_map.items():
            for i in range(len(node_list)):
                for j in range(i + 1, len(node_list)):
                    if node_list[i] in node_ids and node_list[j] in node_ids:
                        key = tuple(sorted([node_list[i], node_list[j]]))
                        if key not in seen_edges:
                            seen_edges.add(key)
                            edges.append({"source": node_list[i], "target": node_list[j], "label": concept})
        link_re = re.compile(r'\[([^\]]*)\]\(([^)]+\.md)\)')
        for md_file in knowledge_path.rglob("*.md"):
            rel = str(md_file.relative_to(knowledge_path)).replace("\\", "/")
            if rel in ("index.md", "log.md", "cross_references.md"):
                continue
            source_id = rel.replace(".md", "")
            if source_id not in node_ids:
                continue
            try:
                content = md_file.read_text(encoding="utf-8")
                for _, link_target in link_re.findall(content):
                    resolved = (md_file.parent / link_target).resolve()
                    try:
                        target_rel = str(resolved.relative_to(knowledge_path)).replace("\\", "/")
                    except ValueError:
                        continue
                    target_id = target_rel.replace(".md", "")
                    if target_id in node_ids and target_id != source_id:
                        key = tuple(sorted([source_id, target_id]))
                        if key not in seen_edges:
                            seen_edges.add(key)
                            edges.append({"source": source_id, "target": target_id, "label": "引用"})
            except Exception:
                pass
        return self._project_graph(
            nodes,
            edges,
            limit=limit,
            focus_id=focus_id,
            depth=depth,
            query=query,
            min_confidence=min_confidence,
        )

    def get_knowledge_graph(
        self,
        book_id: str = "",
        limit: int = 140,
        focus_id: str = "",
        depth: int = 1,
        query: str = "",
        min_confidence: float = 0.0,
    ) -> dict:
        return self.build_knowledge_graph(
            book_id=book_id,
            limit=limit,
            focus_id=focus_id,
            depth=depth,
            query=query,
            min_confidence=min_confidence,
        )

    def _project_graph(
        self,
        nodes: list,
        edges: list,
        limit: int = 140,
        focus_id: str = "",
        depth: int = 1,
        query: str = "",
        min_confidence: float = 0.0,
        include_links: bool = False,
    ) -> dict:
        """Return a small graph projection suitable for interactive UI rendering."""
        limit = max(20, min(int(limit or 140), 500))
        depth = max(1, min(int(depth or 1), 3))
        node_map = {str(n.get("id", "")): n for n in (nodes or []) if n.get("id")}
        filtered_edges = []
        for edge in edges or []:
            source = str(edge.get("source", ""))
            target = str(edge.get("target", ""))
            if source not in node_map or target not in node_map or source == target:
                continue
            try:
                confidence = float(edge.get("confidence", 1.0))
            except Exception:
                confidence = 1.0
            if confidence < min_confidence:
                continue
            filtered_edges.append(edge)

        adjacency = {}
        degree = {node_id: 0 for node_id in node_map}
        for edge in filtered_edges:
            source = str(edge.get("source", ""))
            target = str(edge.get("target", ""))
            adjacency.setdefault(source, set()).add(target)
            adjacency.setdefault(target, set()).add(source)
            degree[source] = degree.get(source, 0) + 1
            degree[target] = degree.get(target, 0) + 1

        selected_ids = set()
        query = (query or "").strip().lower()
        if focus_id and focus_id in node_map:
            frontier = {focus_id}
            selected_ids.add(focus_id)
            for _ in range(depth):
                next_frontier = set()
                for node_id in frontier:
                    next_frontier.update(adjacency.get(node_id, set()))
                selected_ids.update(next_frontier)
                frontier = next_frontier
        elif query:
            for node_id, node in node_map.items():
                haystack = " ".join(
                    str(node.get(k, "")) for k in ("id", "label", "name", "category", "summary")
                ).lower()
                if query in haystack:
                    selected_ids.add(node_id)
                    selected_ids.update(adjacency.get(node_id, set()))
        else:
            ranked = sorted(
                node_map.keys(),
                key=lambda node_id: (
                    degree.get(node_id, 0),
                    1 if str(node_map[node_id].get("category", "")).lower() in ("page", "chunk") else 0,
                    str(node_map[node_id].get("label", node_id)),
                ),
                reverse=True,
            )
            selected_ids.update(ranked[:limit])

        if len(selected_ids) > limit:
            selected_ids = set(
                sorted(
                    selected_ids,
                    key=lambda node_id: (degree.get(node_id, 0), str(node_map[node_id].get("label", node_id))),
                    reverse=True,
                )[:limit]
            )

        projected_edges = [
            edge for edge in filtered_edges
            if str(edge.get("source", "")) in selected_ids and str(edge.get("target", "")) in selected_ids
        ]
        projected_nodes = []
        for node_id in selected_ids:
            node = dict(node_map[node_id])
            node["degree"] = degree.get(node_id, 0)
            if len(str(node.get("summary", ""))) > 240:
                node["summary"] = str(node.get("summary", ""))[:240].rstrip() + "..."
            projected_nodes.append(node)
        projected_nodes.sort(key=lambda n: (n.get("degree", 0), str(n.get("label", n.get("id", "")))), reverse=True)

        category_counts = {}
        for node in nodes or []:
            category = node.get("category") or node.get("type") or "node"
            category_counts[category] = category_counts.get(category, 0) + 1
        clusters = [
            {"category": category, "count": count}
            for category, count in sorted(category_counts.items(), key=lambda item: item[1], reverse=True)
        ]

        result = {
            "nodes": projected_nodes,
            "edges": projected_edges,
            "total_nodes": len(node_map),
            "total_edges": len(filtered_edges),
            "returned_nodes": len(projected_nodes),
            "returned_edges": len(projected_edges),
            "has_more": len(node_map) > len(projected_nodes),
            "mode": "focus" if focus_id else ("search" if query else "summary"),
            "focus_id": focus_id,
            "clusters": clusters[:20],
        }
        if include_links:
            result["links"] = projected_edges
        return result

    def link_to_skill(self, skill_name: str, source_paths: list) -> dict:
        if not skill_name or ".." in skill_name or "/" in skill_name or "\\" in skill_name:
            raise ValueError("invalid skill name")
        for sp in (source_paths or []):
            if ".." in sp:
                raise ValueError(f"invalid source path: {sp}")
        workspace_root = self.workspace_root
        custom_skills_dir = os.path.join(workspace_root, "skills")
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        builtin_skills_dir = os.path.join(project_root, "skills")
        skill_dir = None
        for base in [custom_skills_dir, builtin_skills_dir]:
            candidate = os.path.join(base, skill_name)
            skill_md = os.path.join(candidate, "SKILL.md")
            if os.path.isfile(skill_md):
                skill_dir = candidate
                break
        if not skill_dir:
            raise FileNotFoundError(f"skill not found: {skill_name}")
        skill_md_path = os.path.join(skill_dir, "SKILL.md")
        with open(skill_md_path, "r", encoding="utf-8") as f:
            content = f.read()
        knowledge_section = "\n## Knowledge Sources\n"
        for sp in source_paths:
            knowledge_section += f"- {sp}\n"
        if "## Knowledge Sources" in content:
            lines = content.split("\n")
            new_lines = []
            in_section = False
            for line in lines:
                if line.strip() == "## Knowledge Sources":
                    in_section = True
                    new_lines.append(line)
                    continue
                if in_section and line.startswith("## ") and line.strip() != "## Knowledge Sources":
                    in_section = False
                if in_section:
                    pass
                new_lines.append(line)
            insert_idx = len(new_lines)
            for i, line in enumerate(new_lines):
                if line.strip() == "## Knowledge Sources":
                    j = i + 1
                    while j < len(new_lines) and not new_lines[j].startswith("## "):
                        j += 1
                    insert_idx = j
                    break
            addition_lines = []
            for sp in source_paths:
                addition_lines.append(f"- {sp}")
            new_lines[insert_idx:insert_idx] = addition_lines
            content = "\n".join(new_lines)
        else:
            content = content.rstrip("\n") + "\n" + knowledge_section
        with open(skill_md_path, "w", encoding="utf-8") as f:
            f.write(content)
        return {"skill_name": skill_name, "linked_sources": source_paths}

    def dispatch(self, action: str, payload: Optional[dict] = None) -> dict:
        payload = payload or {}
        try:
            if action == "list":
                book_id = payload.get("book_id", "")
                result = self.list_tree(book_id=book_id)
                return {"action": action, "code": 200, "message": "success", "payload": result}
            elif action == "read":
                path = payload.get("path")
                book_id = payload.get("book_id", "")
                if not path:
                    return {"action": action, "code": 400, "message": "path is required", "payload": None}
                result = self.read_file(path, book_id=book_id)
                return {"action": action, "code": 200, "message": "success", "payload": result}
            elif action == "graph":
                book_id = payload.get("book_id", "")
                result = self.build_graph(book_id=book_id)
                return {"action": action, "code": 200, "message": "success", "payload": result}
            elif action == "organize":
                book_id = payload.get("book_id", "")
                result = self.organize_knowledge(book_id=book_id)
                return {"action": action, "code": 200, "message": "success", "payload": result}
            elif action == "knowledge_graph":
                book_id = payload.get("book_id", "")
                result = self.get_knowledge_graph(book_id=book_id)
                return {"action": action, "code": 200, "message": "success", "payload": result}
            else:
                return {"action": action, "code": 400, "message": f"unknown action: {action}", "payload": None}
        except ValueError as e:
            return {"action": action, "code": 403, "message": str(e), "payload": None}
        except FileNotFoundError as e:
            return {"action": action, "code": 404, "message": str(e), "payload": None}
        except Exception as e:
            logger.error(f"[KnowledgeService] dispatch error: action={action}, error={e}")
            return {"action": action, "code": 500, "message": str(e), "payload": None}
