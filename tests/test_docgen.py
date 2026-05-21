import sys
import os
import tempfile
from docx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from agent.textbook.docgen.style_template import StyleTemplate, get_template, TEMPLATES
from agent.textbook.docgen.word_generator import MarkdownToWordConverter
from agent.textbook.docgen.image_handler import ImageHandler


def test_style_templates():
    academic = get_template("academic")
    assert academic.name == "academic"
    assert academic.heading1_font == "黑体"
    assert academic.heading1_size == 16
    assert academic.body_font == "仿宋"
    assert academic.body_size == 12

    modern = get_template("modern")
    assert modern.name == "modern"
    assert modern.heading1_font == "微软雅黑"
    assert modern.heading1_size == 24
    assert modern.body_font == "微软雅黑"
    assert modern.body_first_line_indent == 0

    training = get_template("training")
    assert training.name == "training"
    assert training.heading1_size == 26
    assert training.body_size == 14
    assert training.body_line_spacing == 2.0

    assert len(TEMPLATES) == 3
    fallback = get_template("nonexistent")
    assert fallback.name == "academic"


def test_converter_heading():
    converter = MarkdownToWordConverter()
    converter.add_heading("第一章 绪论", 1)
    converter.add_heading("1.1 研究背景", 2)
    converter.add_heading("1.1.1 国内现状", 3)

    headings = [p for p in converter.doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) == 3
    assert headings[0].text == "第一章 绪论"
    assert headings[1].text == "1.1 研究背景"
    assert headings[2].text == "1.1.1 国内现状"

    converter.add_heading("无效级别", 5)
    last_heading = [p for p in converter.doc.paragraphs if p.style.name.startswith("Heading")][-1]
    assert last_heading.style.name == "Heading 3"


def test_converter_paragraph():
    converter = MarkdownToWordConverter()
    p = converter.add_paragraph("这是**粗体**和*斜体*以及`代码`的混合文本")

    assert len(p.runs) >= 3

    bold_run = None
    italic_run = None
    code_run = None
    for run in p.runs:
        if run.text == "粗体":
            bold_run = run
        elif run.text == "斜体":
            italic_run = run
        elif run.text == "代码":
            code_run = run

    assert bold_run is not None
    assert bold_run.bold is True
    assert italic_run is not None
    assert italic_run.italic is True
    assert code_run is not None
    assert code_run.font.name == "Times New Roman"

    assert p.paragraph_format.line_spacing == converter.template.body_line_spacing


def test_converter_code_block():
    converter = MarkdownToWordConverter()
    code = "def hello():\n    print('Hello, World!')"
    converter.add_code_block(code, "python")

    last_p = converter.doc.paragraphs[-1]
    assert last_p.text == code
    assert last_p.runs[0].font.name == "Times New Roman"
    assert last_p.paragraph_format.left_indent is not None


def test_converter_table():
    converter = MarkdownToWordConverter()
    headers = ["姓名", "年龄", "城市"]
    rows = [
        ["张三", "25", "北京"],
        ["李四", "30", "上海"],
    ]
    converter.add_table(headers, rows)

    table = converter.doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert table.rows[0].cells[0].text == "姓名"
    assert table.rows[1].cells[0].text == "张三"
    assert table.rows[2].cells[2].text == "上海"

    for run in table.rows[0].cells[0].paragraphs[0].runs:
        assert run.bold is True


def test_converter_export_style_requirements():
    converter = MarkdownToWordConverter()
    converter.convert_markdown(
        "# 一级标题 A1\n\n"
        "正文 ABC 123\n\n"
        "### 三级标题 B2\n\n"
        "表1 示例表\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n"
        "注：这是注释。"
    )

    h1 = next(p for p in converter.doc.paragraphs if p.text == "一级标题 A1")
    assert h1.paragraph_format.line_spacing == 1.5
    assert h1.paragraph_format.space_before.pt == 8
    assert h1.paragraph_format.space_after.pt == 8
    assert h1.runs[0].font.size.pt == 16
    assert h1.runs[0].bold is True
    assert h1.runs[0]._element.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert h1.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"

    body = next(p for p in converter.doc.paragraphs if p.text == "正文 ABC 123")
    assert body.paragraph_format.line_spacing == 1.5
    assert body.runs[0].font.size.pt == 12
    assert body.runs[0]._element.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert body.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "仿宋"

    h3 = next(p for p in converter.doc.paragraphs if p.text == "三级标题 B2")
    assert h3.runs[0].bold in (False, None)
    assert h3.runs[0].font.size.pt == 16

    note = next(p for p in converter.doc.paragraphs if p.text.startswith("注："))
    assert note.runs[0].font.size.pt == 10.5
    assert note.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"

    table_para = converter.doc.tables[0].rows[1].cells[0].paragraphs[0]
    assert table_para.paragraph_format.first_line_indent.pt == 0
    assert table_para.paragraph_format.alignment == 1
    assert table_para.paragraph_format.line_spacing == 1.0
    assert table_para.runs[0].font.size.pt == 10.5
    assert table_para.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "仿宋"


def test_converter_list():
    converter = MarkdownToWordConverter()

    bullet_items = ["第一项", "第二项", "第三项"]
    converter.add_bullet_list(bullet_items)

    numbered_items = ["步骤一", "步骤二", "步骤三"]
    converter.add_numbered_list(numbered_items)

    bullet_paras = [p for p in converter.doc.paragraphs if p.style.name == 'List Bullet']
    assert len(bullet_paras) == 3
    assert bullet_paras[0].text == "第一项"

    numbered_paras = [p for p in converter.doc.paragraphs if p.style.name == 'List Number']
    assert len(numbered_paras) == 3
    assert numbered_paras[0].text == "步骤一"


def test_converter_full_markdown():
    markdown_text = """# 第一章 Python基础

## 1.1 变量与类型

Python是**动态类型**语言，支持*多种数据类型*。

使用`type()`函数查看类型：

```python
x = 42
print(type(x))
```

### 1.1.1 数值类型

- 整数 int
- 浮点数 float
- 复数 complex

1. 第一步
2. 第二步
3. 第三步

> 引用文本示例

| 类型 | 示例 | 说明 |
|------|------|------|
| int | 42 | 整数 |
| float | 3.14 | 浮点数 |

普通段落文本。
"""

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "test_output.docx")
        result_path = MarkdownToWordConverter.from_markdown(
            markdown_text, template_name="academic", output_path=output_path
        )
        assert result_path == output_path
        assert os.path.exists(output_path)
        assert os.path.getsize(output_path) > 0

        converter = MarkdownToWordConverter(get_template("academic"))
        converter.convert_markdown(markdown_text)
        assert len(converter.doc.paragraphs) > 0

        heading_texts = [p.text for p in converter.doc.paragraphs if p.style.name.startswith("Heading")]
        assert "第一章 Python基础" in heading_texts
        assert "1.1 变量与类型" in heading_texts
        assert "1.1.1 数值类型" in heading_texts

        assert len(converter.doc.tables) == 1
        table = converter.doc.tables[0]
        assert table.rows[0].cells[0].text == "类型"


def test_image_handler_resize():
    with tempfile.TemporaryDirectory() as tmpdir:
        from PIL import Image

        img = Image.new("RGB", (1200, 900), color=(255, 0, 0))
        input_path = os.path.join(tmpdir, "input.png")
        output_path = os.path.join(tmpdir, "output.png")
        img.save(input_path)

        ImageHandler.resize_image(input_path, output_path, max_width=800, max_height=600)

        assert os.path.exists(output_path)
        with Image.open(output_path) as resized:
            w, h = resized.size
            assert w <= 800
            assert h <= 600

        dims = ImageHandler.get_image_dimensions(input_path)
        assert dims == (1200, 900)

        convert_path = os.path.join(tmpdir, "converted.jpg")
        ImageHandler.convert_format(input_path, convert_path, format="JPEG")
        assert os.path.exists(convert_path)
